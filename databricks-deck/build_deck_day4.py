#!/usr/bin/env python3
"""
build_deck_day4.py
------------------
DAY 4 of the Databricks training series — DATA PROCESSING CONCEPTS.
How modern data flows through an enterprise platform from ingestion to
analytics. Same design system & layout as Days 1-3, infographic-rich.

Run:  python3 build_deck_day4.py
Out:  Databricks-Day4-Data-Processing-Concepts.pptx
"""

from math import ceil
from deck_kit import *
from build_deck import (flow_steps, grid_cards, pipeline, timeline,
                        vs_table, dark_bg, takeaway, s_divider)
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "Databricks-Day4-Data-Processing-Concepts.pptx"
PARTS = ["Sources", "Ingestion", "Transform", "Processing", "Platform"]


def progress(s, active):
    x = SW - MX - 5.0
    for i, lab in enumerate(PARTS):
        on = (i == active)
        cxp = x + i * 1.18
        circle(s, cxp, 1.46, 0.12 if on else 0.085, LAVA if on else OAT_3)
        text(s, cxp - 0.55, 1.55, 1.1, 0.22, lab, size=7.2,
             color=(NAVY if on else GRAY_LT), align=PP_ALIGN.CENTER, bold=on, font=FONT_SB)


def note_band(s, y, runs, h=1.0, fill=OAT_2):
    card(s, MX, y, CW, h, fill=fill, line=LINE, shadow=False)
    text(s, MX + 0.38, y, CW - 0.76, h, runs, anchor=MSO_ANCHOR.MIDDLE, spacing=1.25)


def journey(s, stages, y=2.65, h=2.5, badge=True, title_size=11, sub_size=8.4):
    """Horizontal numbered icon-stage flow (no example cards)."""
    n = len(stages)
    aw = 0.16
    cw = (CW - (n - 1) * aw) / n
    for i, st in enumerate(stages):
        x = MX + i * (cw + aw)
        card(s, x, y, cw, h, fill=WHITE, line=LINE, shadow=True)
        band(s, x, y, cw, 0.1, st["col"])
        icon_tile(s, x + cw / 2, y + 0.7, 0.8, st["icon"], st["col"], shadow=False)
        text(s, x + 0.04, y + 1.22, cw - 0.08, 0.42, st["t"], size=title_size, color=NAVY,
             bold=True, align=PP_ALIGN.CENTER, spacing=0.9)
        text(s, x + 0.04, y + 1.66, cw - 0.08, 0.8, st["sub"], size=sub_size, color=GRAY_2,
             align=PP_ALIGN.CENTER, spacing=1.0)
        if badge:
            number_badge(s, x + 0.26, y + 0.26, 0.34, i + 1, fill=st["col"], size=10)
        if i < n - 1:
            arrow_h(s, x + cw + 0.0, y + h / 2, aw, GRAY_LT, h=0.16)


# ===========================================================================
#  INTRODUCTION
# ===========================================================================
def cover(prs):
    s = slide(prs, bg=NAVY)
    dark_bg(s)
    # "data flowing through a pipeline" motif on the right
    for i, d in enumerate([0.16, 0.22, 0.28, 0.22, 0.16]):
        circle(s, 9.9 + i * 0.62, 2.7, d, LAVA if i == 2 else NAVY_3)
    shape(s, MSO_SHAPE.RIGHT_ARROW, 9.7, 3.2, 3.0, 0.34, fill=NAVY_2, line=None)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 1.45, 0.09, 2.0, fill=LAVA, line=None, adj=0.5)
    text(s, MX + 0.30, 1.40, 9.5, 0.3, "DATABRICKS ENABLEMENT  ·  DAY 4",
         size=12.5, color=LAVA_LT, bold=True, font=FONT_SB)
    text(s, MX + 0.27, 1.82, 11.4, 1.2, "Data Processing Concepts", size=46, color=WHITE, bold=True)
    text(s, MX + 0.29, 2.86, 11.0, 0.8, "From raw signals to enterprise insight",
         size=27, color=LAVA, bold=True, font=FONT_LT)
    text(s, MX + 0.30, 3.78, 8.6, 0.7,
         "How modern data flows through a platform — sources, ingestion, transformation, "
         "distributed processing, pipelines, quality and storage.",
         size=14, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.25)
    chips = [("01", "Sources & Types"), ("02", "Ingestion"), ("03", "Transform · ETL/ELT"),
             ("04", "Processing & Pipelines"), ("05", "Quality · Storage")]
    cx = MX + 0.30
    for no, label in chips:
        w = 2.16
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx, 4.85, w, 0.58, fill=None,
              line=LINE_DK, line_w=1.2, adj=0.5)
        text(s, cx + 0.14, 4.85, w - 0.2, 0.58,
             [[{"t": no + " ", "color": LAVA, "size": 10.5, "bold": True},
               {"t": label, "color": WHITE, "size": 9.3, "bold": True}]],
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        cx += w + 0.15
    band(s, MX + 0.30, 6.05, 6.5, 0.02, LINE_DK)
    text(s, MX + 0.30, 6.22, 11, 0.4,
         "Audience:  Data & ETL Engineers · Analysts · Scientists · Cloud Engineers · Architects",
         size=11.5, color=GRAY_LT)
    text(s, MX + 0.30, 6.62, 11, 0.4,
         "Beginner → intermediate  ·  vendor-neutral concepts, shown in a Databricks context",
         size=11.5, color=GRAY_LT, italic=True)
    notes(s,
          "Welcome to Day 4 — the conceptual backbone that ties the series together. Where earlier days "
          "covered specific technologies (Spark, the Workspace, Delta Lake), today is about the universal "
          "CONCEPTS of data processing: how raw data becomes trusted insight. We'll follow data left to right "
          "through five parts — sources & types, ingestion, transformation and ETL/ELT, distributed "
          "processing & pipelines, and finally quality & storage — landing on the end-to-end Databricks flow. "
          "The concepts are vendor-neutral, but we'll always anchor them in how Databricks implements them. "
          "Aimed at beginner-to-intermediate practitioners.")


def why_matters(prs):
    s = slide(prs)
    header(s, "Introduction", "Why Data Processing Matters")
    # Raw -> Insight -> Value flow
    steps = [("Raw Data", "noisy, scattered, unusable", "database", GRAY_2),
             ("Insights", "clean, modelled, trusted", "gauge", BLUE),
             ("Business Value", "decisions & outcomes", "target", LAVA)]
    bw = 3.4
    for i, (t, d, ic, col) in enumerate(steps):
        x = MX + i * (bw + 0.65)
        card(s, x, 2.1, bw, 1.7, fill=WHITE, line=LINE, shadow=True)
        icon_tile(s, x + 0.75, 2.95, 0.95, ic, col, shadow=True)
        text(s, x + 1.35, 2.4, bw - 1.5, 0.45, t, size=14.5, color=NAVY, bold=True)
        text(s, x + 1.35, 2.9, bw - 1.5, 0.7, d, size=10.5, color=GRAY_2, spacing=1.05)
        if i < 2:
            arrow_h(s, x + bw + 0.06, 2.95, 0.55, LAVA, h=0.24)
    items = [
        dict(icon="bolt", color=LAVA, title="Faster Decisions",
             desc="Timely, trustworthy data shortens the path from event to action."),
        dict(icon="transaction", color=BLUE, title="Automation",
             desc="Processed data powers automated workflows & operational systems."),
        dict(icon="target", color=GREEN, title="Innovation",
             desc="Clean data fuels ML, AI and new data-driven products."),
        dict(icon="people", color=PURPLE, title="Customer Experience",
             desc="Personalization & real-time service depend on processed data."),
    ]
    grid_cards(s, items, cols=4, x=MX, y=4.15, w=CW, h=2.3, gx=0.25,
               icon_d=0.7, title_size=12.5, desc_size=9.8)
    footer(s)
    notes(s,
          "Frame the 'why' before the 'how'. Raw data on its own is noise — scattered across systems, "
          "inconsistent, and unusable for decisions. Processing is the act of turning that raw data into "
          "insight, and insight into business value. Use the three-step flow to make it concrete, then the "
          "four outcome cards to show why organizations invest: faster decisions, automation, innovation "
          "(ML/AI), and better customer experience. The takeaway: data processing isn't back-office plumbing — "
          "it's the engine that converts data into competitive advantage.")


def data_journey(prs):
    s = slide(prs)
    header(s, "Introduction", "The Data Journey — Seven Stages")
    stages = [
        dict(t="Generate", sub="apps · sensors\nevents", icon="iot", col=NAVY_2),
        dict(t="Collect", sub="ingest from\nsources", icon="funnel", col=TEAL),
        dict(t="Process", sub="clean &\ntransform", icon="gear", col=LAVA),
        dict(t="Store", sub="lake / lakehouse\ntables", icon="database", col=BRONZE),
        dict(t="Analyze", sub="SQL · ML\nstatistics", icon="gauge", col=BLUE),
        dict(t="Visualize", sub="dashboards\n& reports", icon="chart", col=PURPLE),
        dict(t="Act", sub="decisions &\nautomation", icon="target", col=GREEN),
    ]
    journey(s, stages, y=2.75, h=2.7, title_size=11, sub_size=8.2)
    note_band(s, 5.75,
              [[{"t": "Every data platform implements this same arc.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "The rest of today zooms into each stage — and shows how Databricks delivers it end to end.",
                 "color": GRAY, "size": 12.5}]], h=0.95)
    footer(s)
    notes(s,
          "This is the spine of the entire deck — the seven-stage data journey that every platform follows. "
          "Data is GENERATED by apps, sensors and events; COLLECTED via ingestion; PROCESSED (cleaned and "
          "transformed); STORED in a lake or lakehouse; ANALYZED with SQL, ML and statistics; VISUALIZED in "
          "dashboards; and finally acted upon — driving decisions and automation. Tell the audience this slide "
          "is the map: each subsequent section zooms into one or more of these stages. Keep it memorable — if "
          "they remember nothing else, they should remember this arc.")


# ===========================================================================
#  PART 1 — SOURCES & TYPES
# ===========================================================================
def sources(prs):
    s = slide(prs)
    header(s, "Data Sources", "Where Enterprise Data Comes From")
    progress(s, 0)
    items = [
        dict(icon="database", color=BLUE, title="Databases", desc="OLTP systems, ERP, CRM — structured records."),
        dict(icon="api", color=LAVA, title="APIs", desc="REST & GraphQL endpoints from internal & SaaS apps."),
        dict(icon="iot", color=TEAL, title="IoT & Sensors", desc="Devices and machines emitting telemetry streams."),
        dict(icon="people", color=PURPLE, title="Applications", desc="Web & mobile apps generating user events."),
        dict(icon="doc", color=GREEN, title="Logs & Files", desc="Server logs, CSV/JSON files, exports & dumps."),
        dict(icon="globe", color=YELLOW, title="Social & Web", desc="Clickstream, social feeds & external data."),
    ]
    grid_cards(s, items, cols=3, x=MX, y=2.0, w=CW, h=4.5, icon_d=0.72,
               title_size=14, desc_size=10.8)
    footer(s)
    notes(s,
          "Enterprises pull data from a sprawl of sources — and the variety is exactly what makes processing "
          "hard. Walk the six categories: transactional DATABASES (the structured core), APIs (internal "
          "services and SaaS like Salesforce), IoT & SENSORS (high-velocity telemetry), APPLICATIONS (web/"
          "mobile event streams), LOGS & FILES (semi-structured exports and machine logs), and SOCIAL & WEB "
          "(clickstream and external feeds). The key teaching point: these sources differ in structure, speed "
          "and volume, so a platform must ingest and reconcile all of them — which sets up the next slide on "
          "data types and the whole ingestion section.")


def data_types(prs):
    s = slide(prs)
    header(s, "Data Sources", "Structured vs. Semi vs. Unstructured")
    progress(s, 0)
    cols = [
        dict(name="Structured", color=BLUE, icon="table",
             rows=[("Examples", "SQL tables, CSV"), ("Schema", "Fixed, defined upfront"),
                   ("Storage", "Warehouses, RDBMS"), ("Share", "~20% of data")]),
        dict(name="Semi-Structured", color=TEAL, icon="doc",
             rows=[("Examples", "JSON, XML, logs"), ("Schema", "Flexible / self-describing"),
                   ("Storage", "Lakes, NoSQL"), ("Share", "Growing fast")]),
        dict(name="Unstructured", color=LAVA, icon="variety",
             rows=[("Examples", "Text, images, video"), ("Schema", "None — raw content"),
                   ("Storage", "Object storage / lake"), ("Share", "~80% of data")]),
    ]
    cw = (CW - 2 * 0.4) / 3
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.4)
        card(s, x, 2.05, cw, 4.35, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.05, cw, 0.78, c["color"])
        icon(s, c["icon"], x + 0.6, 2.44, 0.44, WHITE)
        text(s, x + 1.05, 2.05, cw - 1.1, 0.78, c["name"], size=14.5, color=WHITE,
             bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.92)
        yy = 3.05
        for k, v in c["rows"]:
            text(s, x + 0.32, yy, cw - 0.6, 0.3, k.upper(), size=8.5, color=c["color"],
                 bold=True, font=FONT_SB)
            text(s, x + 0.32, yy + 0.26, cw - 0.6, 0.45, v, size=11.5, color=GRAY, spacing=1.0)
            band(s, x + 0.32, yy + 0.74, cw - 0.64, 0.01, OAT_2)
            yy += 0.82
    footer(s)
    notes(s,
          "Data comes in three shapes, and the mix has shifted dramatically. STRUCTURED data fits neat rows "
          "and columns with a fixed schema — SQL tables, CSVs — and lives in warehouses; it's only about 20% "
          "of enterprise data today. SEMI-STRUCTURED data (JSON, XML, logs) carries its own flexible schema "
          "and is growing fast. UNSTRUCTURED data — text, images, audio, video — has no schema at all and is "
          "now roughly 80% of all data. The implication: traditional warehouses, built only for structured "
          "data, can't handle the modern mix — which is why lakes and lakehouses exist. (Percentages are "
          "common industry estimates.)")


# ===========================================================================
#  PART 2 — INGESTION
# ===========================================================================
def what_ingestion(prs):
    s = slide(prs)
    header(s, "Data Ingestion", "What Is Data Ingestion?")
    progress(s, 1)
    stages = [
        dict(title="Sources", icon="globe", color=NAVY_2,
             examples=["Databases", "APIs · IoT", "Files · logs"]),
        dict(title="Ingestion Layer", icon="funnel", color=LAVA,
             examples=["Connect", "Extract", "Land raw data"]),
        dict(title="Storage", icon="database", color=BLUE,
             examples=["Data lake", "Bronze tables", "Object storage"]),
    ]
    pipeline(s, stages, y_top=2.5, box_h=1.05, ex_h=1.4, label_size=12.5, ex_size=10)
    note_band(s, 5.4,
              [[{"t": "Ingestion is step one of every pipeline.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "It connects to source systems and moves their data into the platform — reliably, and "
                      "without losing or duplicating records.", "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Define ingestion simply: it's the process of getting data FROM source systems INTO your data "
          "platform. It's the first stage of every pipeline and the foundation everything else builds on. The "
          "ingestion layer connects to each source, extracts the data, and lands it (usually raw, in a Bronze "
          "layer). The hard parts — which the next slides cover — are doing this reliably (no lost or "
          "duplicated records) and choosing the right mode: batch (periodic) or streaming (continuous). On "
          "Databricks, Auto Loader and Lakeflow Connect handle both.")


def batch_ingestion(prs):
    s = slide(prs)
    header(s, "Data Ingestion", "Batch Ingestion")
    progress(s, 1)
    events = [
        dict(year="02:00", title="Nightly Load", desc="Yesterday's data ingested in bulk", color=BLUE),
        dict(year="06:00", title="Morning Batch", desc="Incremental files picked up", color=TEAL),
        dict(year="12:00", title="Midday Run", desc="Scheduled API & DB extract", color=LAVA),
        dict(year="18:00", title="Evening Batch", desc="End-of-day reconciliation", color=PURPLE),
    ]
    timeline(s, events, y_center=4.0)
    note_band(s, 5.6,
              [[{"t": "Batch = data ingested on a schedule, in chunks.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Simple, cost-efficient and ideal for large historical loads where minutes or hours of "
                      "latency are acceptable.", "color": GRAY, "size": 12.5}]], h=0.9)
    footer(s)
    notes(s,
          "Batch ingestion loads data in scheduled chunks — for example, every night at 2am, or every few "
          "hours. Use the clock timeline to make the 'periodic' nature visceral. Batch is the workhorse of "
          "data engineering: it's simple, cheap, easy to reason about, and perfect for large historical loads "
          "and reports where some latency (minutes to hours) is fine. The trade-off is freshness — you only "
          "see data as recent as the last run. That limitation is exactly what streaming ingestion, on the "
          "next slide, solves. On Databricks, scheduled Jobs/Workflows and Auto Loader handle incremental "
          "batch efficiently.")


def streaming_ingestion(prs):
    s = slide(prs)
    header(s, "Data Ingestion", "Streaming Ingestion")
    progress(s, 1)
    # sources -> stream bus -> platform
    srcs = [("Kafka", "stream"), ("IoT Devices", "iot"), ("Event Streams", "transaction")]
    for i, (t, ic) in enumerate(srcs):
        y = 2.3 + i * 1.4
        card(s, MX, y, 3.0, 1.15, fill=WHITE, line=LINE, shadow=True)
        icon_tile(s, MX + 0.6, y + 0.57, 0.66, ic, LAVA, shadow=False)
        text(s, MX + 1.1, y, 1.9, 1.15, t, size=12.5, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(MX + 3.0), Inches(y + 0.57),
                                    Inches(5.5), Inches(4.0))
        ln.line.color.rgb = GRAY_LT; ln.line.width = Pt(1.8)
    # continuous stream hub
    circle(s, 6.4, 4.0, 1.6, LAVA, line=WHITE, line_w=3, shadow=True)
    icon(s, "stream", 6.4, 3.72, 0.55, WHITE)
    text(s, 5.6, 4.05, 1.6, 0.6, "Continuous\ningest", size=12, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
    # arrow to platform
    a = shape(s, MSO_SHAPE.RIGHT_ARROW, 7.9, 3.82, 1.2, 0.36, fill=GRAY_LT, line=None)
    card(s, 9.3, 3.2, CW + MX - 9.3, 1.6, fill=NAVY, line=None, shadow=True)
    icon_tile(s, 9.3 + 0.7, 4.0, 0.8, "database", GREEN, shadow=False)
    text(s, 9.3 + 1.25, 3.45, 2.4, 0.5, "Delta Lake", size=14, color=WHITE, bold=True)
    text(s, 9.3 + 1.25, 3.9, 2.5, 0.7, "Real-time Bronze table, processed within seconds",
         size=9.5, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.05)
    note_band(s, 5.6,
              [[{"t": "Streaming = data ingested continuously, in real time.  ", "color": NAVY, "size": 12, "bold": True},
                {"t": "Events flow in as they happen — powering fraud detection, monitoring and live dashboards.",
                 "color": GRAY, "size": 12}]], h=0.9)
    footer(s)
    notes(s,
          "Streaming ingestion processes data continuously, as it arrives, rather than waiting for a schedule. "
          "Sources like Kafka, IoT devices and event streams feed a continuous ingestion layer that lands data "
          "in (for example) a Delta Bronze table within seconds. Use cases demand freshness: fraud detection, "
          "real-time monitoring, live operational dashboards, recommendations. The trade-off vs batch is more "
          "complexity and (often) cost. The good news on Databricks: Structured Streaming and Auto Loader let "
          "you write streaming ingestion with almost the same code as batch, and Delta tables safely accept "
          "continuous writes — so the gap between batch and streaming is small in practice.")


def batch_vs_streaming(prs):
    s = slide(prs)
    header(s, "Data Ingestion", "Batch vs. Streaming")
    progress(s, 1)
    rows = [
        ("Timing", "Scheduled, periodic", "Continuous, real-time"),
        ("Latency", "Minutes to hours", "Seconds or less"),
        ("Data Size", "Large bounded chunks", "Unbounded event flow"),
        ("Complexity", "Simpler to build & operate", "More moving parts"),
        ("Cost", "Lower, predictable", "Higher, always-on"),
        ("Best For", "Reports, history, ETL", "Fraud, IoT, live dashboards"),
    ]
    vs_table(s, MX, 2.0, CW, ("Aspect", "Batch", "Streaming"), rows, hcolors=(BLUE, LAVA))
    footer(s)
    notes(s,
          "Make the choice concrete with a side-by-side. BATCH runs on a schedule with minutes-to-hours "
          "latency, handles large bounded chunks, is simpler and cheaper, and suits reports, historical loads "
          "and classic ETL. STREAMING runs continuously with sub-second latency over an unbounded flow of "
          "events, has more moving parts and higher always-on cost, and suits fraud detection, IoT and live "
          "dashboards. The modern reality: it's not either/or — many pipelines do both, and Delta Lake plus "
          "Structured Streaming let one table and nearly one codebase serve both, so teams pick the right mode "
          "per use case rather than maintaining two separate stacks.")


# ===========================================================================
#  PART 3 — TRANSFORMATION, ETL & ELT
# ===========================================================================
def what_transformation(prs):
    s = slide(prs)
    header(s, "Transformation", "What Is Data Transformation?")
    progress(s, 2)
    steps = [
        dict(title="Raw Data", icon="database", color=NAVY_2, badge=1,
             desc="As-ingested — messy, mixed, untrusted."),
        dict(title="Clean", icon="filter", color=TEAL, badge=2,
             desc="Fix nulls, types, errors & duplicates."),
        dict(title="Standardize", icon="sort", color=BLUE, badge=3,
             desc="Conform formats, units & keys."),
        dict(title="Enrich", icon="merge", color=PURPLE, badge=4,
             desc="Join reference data & derive fields."),
        dict(title="Processed", icon="check", color=GREEN, badge=5,
             desc="Trusted, analytics-ready data."),
    ]
    flow_steps(s, steps, y_top=2.35, card_h=2.5, icon_d=1.0, arrow=True)
    note_band(s, 5.25,
              [[{"t": "Transformation turns raw data into trustworthy data.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "It's the heart of processing — and maps directly onto the Bronze → Silver → Gold "
                      "medallion layers.", "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Transformation is where raw data becomes useful — the core of 'data processing.' Walk the arc: RAW "
          "data is messy and untrusted; CLEAN fixes nulls, wrong types, errors and duplicates; STANDARDIZE "
          "conforms formats, units and keys so data from different sources lines up; ENRICH joins reference "
          "data and derives new fields; and the result is PROCESSED, analytics-ready data. Crucially, this is "
          "exactly the medallion pattern from earlier days — Bronze (raw) → Silver (clean/standardized) → "
          "Gold (enriched/aggregated). On Databricks this runs as Spark/SQL transformations over Delta tables.")


def common_transformations(prs):
    s = slide(prs)
    header(s, "Transformation", "Common Transformations")
    progress(s, 2)
    items = [
        dict(icon="filter", color=LAVA, title="Filtering", desc="Keep only the rows that matter (WHERE)."),
        dict(icon="merge", color=BLUE, title="Aggregation", desc="Summarize with sums, counts & averages."),
        dict(icon="link", color=TEAL, title="Joining", desc="Combine datasets on shared keys."),
        dict(icon="sort", color=PURPLE, title="Sorting", desc="Order data by one or more columns."),
        dict(icon="variety", color=YELLOW, title="Deduplication", desc="Remove duplicate records safely."),
        dict(icon="shield", color=GREEN, title="Validation", desc="Enforce rules, types & quality checks."),
    ]
    grid_cards(s, items, cols=3, x=MX, y=2.0, w=CW, h=4.5, icon_d=0.72,
               title_size=14, desc_size=10.8)
    footer(s)
    notes(s,
          "These are the building blocks engineers combine in every pipeline. FILTERING keeps relevant rows. "
          "AGGREGATION rolls data up into sums, counts and averages (the basis of most reporting). JOINING "
          "combines datasets on keys — the workhorse of enrichment. SORTING orders data. DEDUPLICATION "
          "removes duplicates (critical after retries or streaming). VALIDATION enforces rules and types to "
          "guarantee quality. In Spark/SQL these are expressed declaratively — filter, groupBy, join, "
          "orderBy, dropDuplicates — and the engine optimizes how they run. Most real transformations are "
          "just compositions of these six.")


def etl(prs):
    s = slide(prs)
    header(s, "ETL & ELT", "ETL — Extract, Transform, Load")
    progress(s, 2)
    steps = [
        dict(title="Extract", icon="funnel", color=BLUE, badge=1,
             desc="Pull data from source systems."),
        dict(title="Transform", icon="gear", color=LAVA, badge=2,
             desc="Clean & reshape on a staging engine — before loading."),
        dict(title="Load", icon="database", color=GREEN, badge=3,
             desc="Write the finished data into the warehouse."),
    ]
    flow_steps(s, steps, y_top=2.4, card_h=2.55, icon_d=1.05, arrow=True)
    note_band(s, 5.3,
              [[{"t": "ETL transforms data BEFORE loading.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "The classic warehouse approach — transformation happens on a separate engine, so only "
                      "clean, modelled data lands in the target.", "color": GRAY, "size": 12.5}]], h=0.95)
    footer(s)
    notes(s,
          "ETL — Extract, Transform, Load — is the classic data-warehouse pattern. You EXTRACT from sources, "
          "TRANSFORM on a separate staging/ETL engine, and only then LOAD the clean, modelled result into the "
          "warehouse. The defining trait: transformation happens BEFORE loading, so the warehouse only ever "
          "holds polished data. This made sense when storage and warehouse compute were expensive — you "
          "didn't want raw data sitting in a costly warehouse. The downside is rigidity and a separate "
          "transformation tier. Contrast this with ELT on the next slide, which the cloud made possible.")


def elt(prs):
    s = slide(prs)
    header(s, "ETL & ELT", "ELT — Extract, Load, Transform")
    progress(s, 2)
    steps = [
        dict(title="Extract", icon="funnel", color=BLUE, badge=1,
             desc="Pull raw data from source systems."),
        dict(title="Load", icon="database", color=GREEN, badge=2,
             desc="Land raw data straight into the lake/lakehouse."),
        dict(title="Transform", icon="gear", color=LAVA, badge=3,
             desc="Transform in-place with scalable cloud compute."),
    ]
    flow_steps(s, steps, y_top=2.4, card_h=2.55, icon_d=1.05, arrow=True)
    note_band(s, 5.3,
              [[{"t": "ELT loads first, then transforms in place.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "The modern cloud approach — cheap storage keeps raw data, and elastic compute "
                      "(Spark) transforms it on demand. This is the Lakehouse / medallion model.", "color": GRAY, "size": 12.5}]],
              h=0.95)
    footer(s)
    notes(s,
          "ELT — Extract, Load, Transform — flips the last two steps and is the modern, cloud-native approach. "
          "You EXTRACT, LOAD raw data straight into cheap lake/lakehouse storage, and TRANSFORM it there using "
          "elastic compute like Spark. Why the flip works now: cloud storage is cheap (so keeping raw data is "
          "fine) and cloud compute is elastic and powerful (so transforming at scale in-place is easy). "
          "Benefits: you keep the raw data (great for reprocessing and ML), transformations are flexible, and "
          "there's no separate ETL tier. This is exactly the medallion pattern — load to Bronze, then "
          "transform to Silver and Gold. Databricks is built for ELT.")


def etl_vs_elt(prs):
    s = slide(prs)
    header(s, "ETL & ELT", "ETL vs. ELT")
    progress(s, 2)
    rows = [
        ("Order", "Transform → Load", "Load → Transform"),
        ("Transform Where", "Separate ETL engine", "In the lake/lakehouse"),
        ("Raw Data Kept", "No — only modelled data", "Yes — full raw history"),
        ("Speed & Scale", "Limited by ETL tier", "Elastic cloud compute"),
        ("Cost", "Higher infra overhead", "Cheap storage + on-demand"),
        ("Cloud-Native", "Legacy warehouse era", "Modern Lakehouse standard"),
    ]
    vs_table(s, MX, 2.0, CW, ("Aspect", "ETL", "ELT"), rows, hcolors=(NAVY_2, LAVA))
    footer(s)
    notes(s,
          "Summarize the shift. The headline difference is ORDER — ETL transforms before loading, ELT loads "
          "then transforms. That cascades into everything else: ETL transforms on a separate engine and keeps "
          "only modelled data; ELT transforms in the lakehouse and keeps the full raw history (valuable for "
          "reprocessing and ML). ELT scales with elastic cloud compute and leans on cheap storage, while ETL "
          "carries more fixed infrastructure overhead. ETL was right for the on-prem warehouse era; ELT is the "
          "modern lakehouse standard. Be balanced — ETL still appears where strict pre-load governance is "
          "required — but the industry has shifted decisively to ELT, which is what Databricks is designed for.")


# ===========================================================================
#  PART 4 — DISTRIBUTED PROCESSING & PIPELINES
# ===========================================================================
def why_distributed(prs):
    s = slide(prs)
    header(s, "Processing", "Why Distributed Processing?")
    progress(s, 3)
    before_after(s, 2.05, 3.5,
                 "Single Server", [
                     "One machine — a hard capacity ceiling",
                     "Scales UP (bigger, pricier box)",
                     "Slows to a crawl at TB+ scale",
                     "A single point of failure"],
                 "Cluster Processing", [
                     "Many machines working in parallel",
                     "Scales OUT — just add nodes",
                     "Handles petabytes by partitioning",
                     "Fault-tolerant — survives node loss"],
                 left_icon="warning", right_icon="cluster", right_color=GREEN_DK)
    note_band(s, 5.85,
              [[{"t": "Big data outgrew the single machine.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Distributed processing splits work across a cluster — the core idea behind Spark and "
                      "every modern data platform.", "color": GRAY, "size": 12.5}]], h=0.85)
    footer(s)
    notes(s,
          "Explain the fundamental shift behind all modern data processing. A SINGLE SERVER has a hard ceiling "
          "— you can only buy so big a machine (scale UP), it's a single point of failure, and at terabyte-"
          "plus scale it grinds to a halt. DISTRIBUTED / CLUSTER processing instead splits the work across "
          "many machines running in parallel (scale OUT) — add nodes for more capacity, partition data to "
          "handle petabytes, and tolerate failures because no single node is critical. This 'scale out, not "
          "up' principle is the foundation of Spark and every cloud data platform. It's why we can process "
          "data volumes that were impossible a decade ago.")


def spark_model(prs):
    s = slide(prs)
    header(s, "Processing", "The Apache Spark Processing Model")
    progress(s, 3)
    steps = [
        dict(title="Driver", icon="gear", color=NAVY_2, badge=1,
             desc="Plans the job and coordinates the cluster."),
        dict(title="Executors", icon="cluster", color=BLUE, badge=2,
             desc="Worker processes that run tasks in parallel."),
        dict(title="Tasks", icon="bolt", color=LAVA, badge=3,
             desc="Units of work — one per data partition."),
    ]
    flow_steps(s, steps, y_top=2.4, card_h=2.55, icon_d=1.05, arrow=True)
    note_band(s, 5.3,
              [[{"t": "Spark is the engine of distributed processing.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "A Driver breaks work into Tasks that Executors run in parallel across the cluster — "
                      "in memory, and fault-tolerant.", "color": GRAY, "size": 12.5}]], h=0.95)
    footer(s)
    notes(s,
          "Introduce (or recap, for those who saw Day 2) Spark's execution model — the canonical distributed "
          "engine. The DRIVER is the brain: it plans the job and coordinates. EXECUTORS are worker processes "
          "spread across the cluster that actually run the computation. TASKS are the atoms of work — each one "
          "processes a single partition of the data, and many run in parallel, which is where the speed comes "
          "from. Keep it crisp: driver coordinates, executors work, tasks run in parallel on partitions, all "
          "in memory. This is how a logical transformation becomes thousands of parallel operations across a "
          "cluster.")


def how_spark_processes(prs):
    s = slide(prs)
    header(s, "Processing", "How Spark Processes Data")
    progress(s, 3)
    steps = [
        dict(title="Read", icon="funnel", color=NAVY_2, badge=1, desc="Load data from the source or table."),
        dict(title="Partition", icon="variety", color=BLUE, badge=2, desc="Split into chunks across the cluster."),
        dict(title="Transform", icon="gear", color=TEAL, badge=3, desc="Apply operations to each partition."),
        dict(title="Execute", icon="bolt", color=LAVA, badge=4, desc="Run tasks in parallel on executors."),
        dict(title="Write", icon="database", color=GREEN, badge=5, desc="Persist results back to storage."),
    ]
    flow_steps(s, steps, y_top=2.35, card_h=2.5, icon_d=1.0, arrow=True)
    note_band(s, 5.25,
              [[{"t": "Partitioning is the secret to scale.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "By splitting data into partitions and processing them in parallel, Spark turns a huge "
                      "job into many small ones that finish fast.", "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Now the data's-eye view of a Spark job. READ loads the input. PARTITION splits it into chunks "
          "distributed across the cluster — this is the key step. TRANSFORM applies your operations to each "
          "partition independently. EXECUTE runs those as parallel tasks on the executors. WRITE persists the "
          "results. The teaching point: partitioning is what makes scale possible — a 1TB job becomes "
          "thousands of small partition-sized tasks that run simultaneously. Mention that Spark is lazy (it "
          "plans the whole DAG before running) and runs in memory, which is why it's fast. This connects the "
          "abstract 'distributed processing' idea to concrete mechanics.")


def what_pipeline(prs):
    s = slide(prs)
    header(s, "Pipelines", "What Is a Data Pipeline?")
    progress(s, 3)
    stages = [
        dict(title="Ingest", icon="funnel", color=TEAL, examples=["Auto Loader", "Connectors"]),
        dict(title="Process", icon="gear", color=LAVA, examples=["Clean", "Transform"]),
        dict(title="Store", icon="database", color=BLUE, examples=["Delta tables", "Medallion"]),
        dict(title="Serve", icon="chart", color=PURPLE, examples=["BI · ML", "Apps"]),
    ]
    pipeline(s, stages, y_top=2.5, box_h=1.0, ex_h=1.25, label_size=12.5, ex_size=10)
    note_band(s, 5.25,
              [[{"t": "A pipeline is an automated, repeatable data flow.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "It moves data through ingest → process → store → serve on a schedule or in real time — "
                      "reliably, without manual steps.", "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "A data pipeline is an automated, repeatable sequence of steps that moves data from source to "
          "destination — ingest, process, store, serve. The keyword is AUTOMATED: instead of someone running "
          "scripts by hand, a pipeline runs on a schedule or continuously, reliably, with error handling and "
          "monitoring. Think of it as a factory assembly line for data. The next two slides break this into a "
          "concrete architecture and the orchestration that runs it. On Databricks, pipelines are built with "
          "Workflows, Delta Live Tables / Lakeflow Declarative Pipelines, and notebooks.")


def pipeline_arch(prs):
    s = slide(prs)
    header(s, "Pipelines", "Modern Data Pipeline Architecture")
    progress(s, 3)
    stages = [
        dict(t="Sources", sub="DBs · APIs\nIoT · files", icon="globe", col=NAVY_2),
        dict(t="Ingestion", sub="Auto Loader\nbatch + stream", icon="funnel", col=TEAL),
        dict(t="Processing", sub="Spark\ntransform", icon="gear", col=LAVA),
        dict(t="Storage", sub="Delta\nmedallion", icon="database", col=BLUE),
        dict(t="Analytics", sub="BI · ML\nGenAI", icon="gauge", col=PURPLE),
    ]
    journey(s, stages, y=2.65, h=2.55, title_size=12, sub_size=8.6)
    gov = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.35, CW, 0.66, fill=NAVY, line=None, adj=0.2)
    icon(s, "lock", MX + 0.48, 5.68, 0.3, LAVA_LT)
    text(s, MX + 0.9, 5.35, CW - 1.1, 0.66,
         [[{"t": "Governance & orchestration", "color": LAVA_LT, "size": 11.5, "bold": True},
           {"t": "  span every stage — security, lineage, scheduling & monitoring.", "color": WHITE, "size": 11}]],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, 6.15, CW, 0.4,
         "Five stages, one governed flow — the blueprint behind every production data platform.",
         size=11.5, color=GRAY_2, italic=True, align=PP_ALIGN.CENTER)
    footer(s)
    notes(s,
          "This is the reference architecture every modern pipeline follows: SOURCES → INGESTION → PROCESSING "
          "→ STORAGE → ANALYTICS, with governance and orchestration spanning all of it. It's the seven-stage "
          "data journey from the intro, compressed into the five engineering stages teams actually build. "
          "Note how it maps to Databricks: Auto Loader for ingestion, Spark for processing, Delta + medallion "
          "for storage, and SQL/BI/ML for analytics — all governed by Unity Catalog and run by Workflows. "
          "This slide is the bridge between the abstract concepts and a real platform.")


def orchestration(prs):
    s = slide(prs)
    header(s, "Pipelines", "Pipeline Orchestration")
    progress(s, 3)
    items = [
        dict(icon="flow", color=LAVA, title="Jobs & Tasks",
             desc="Define multi-step workflows with dependencies between tasks."),
        dict(icon="clock", color=BLUE, title="Scheduling",
             desc="Run on a schedule, on a trigger, or continuously."),
        dict(icon="branch", color=TEAL, title="Dependencies",
             desc="Tasks run in the right order; failures stop downstream steps."),
        dict(icon="gauge", color=GREEN, title="Monitoring",
             desc="Track runs, get alerts, retry failures & view lineage."),
    ]
    grid_cards(s, items, cols=2, x=MX, y=2.0, w=CW, h=4.45, gx=0.3, gy=0.3,
               icon_d=0.78, title_size=14.5, desc_size=11)
    footer(s)
    notes(s,
          "Orchestration is the conductor that makes a pipeline production-grade. JOBS & TASKS define the "
          "multi-step workflow — a DAG of work. SCHEDULING decides when it runs (cron, triggers, or "
          "continuous). DEPENDENCIES ensure tasks run in the correct order and that a failure halts the "
          "downstream steps instead of producing bad data. MONITORING tracks every run, sends alerts, retries "
          "transient failures, and surfaces lineage. Without orchestration you have scripts; with it you have "
          "a reliable, observable pipeline. On Databricks this is Workflows (orchestration) plus Lakeflow / "
          "Delta Live Tables for declarative pipelines. Tools like Airflow play the same role elsewhere.")


# ===========================================================================
#  PART 5 — QUALITY, STORAGE & DATABRICKS
# ===========================================================================
def why_quality(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "Why Data Quality Matters")
    progress(s, 4)
    items = [
        dict(icon="coins", color=LAVA, title="Costly Mistakes",
             desc="Bad data drives wrong decisions — and real financial loss."),
        dict(icon="warning", color=YELLOW, title="Broken Trust",
             desc="Once stakeholders distrust a dashboard, adoption collapses."),
        dict(icon="clock", color=BLUE, title="Wasted Time",
             desc="Teams burn cycles reconciling and re-checking numbers."),
        dict(icon="lock", color=PURPLE, title="Compliance Risk",
             desc="Inaccurate or unauditable data creates regulatory exposure."),
    ]
    grid_cards(s, items, cols=2, x=MX, y=2.0, w=7.5, h=4.45, gx=0.3, gy=0.3,
               icon_d=0.8, title_size=14.5, desc_size=11)
    card(s, 8.3, 2.0, CW + MX - 8.3, 4.45, fill=NAVY, line=None, shadow=True)
    text(s, 8.65, 2.3, 4.0, 0.3, "THE PRINCIPLE", size=11, color=LAVA_LT, bold=True, font=FONT_SB)
    icon_tile(s, 8.3 + (CW + MX - 8.3) / 2, 3.5, 1.45, "shield", LAVA, shadow=True)
    text(s, 8.5, 4.5, CW + MX - 8.3 - 0.4, 0.9,
         ["Garbage in,", "garbage out."], size=18, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, spacing=1.05)
    text(s, 8.5, 5.55, CW + MX - 8.3 - 0.4, 0.7,
         "No amount of processing fixes data that was wrong at the source.",
         size=10.5, color=RGBColor(0xC4, 0xD0, 0xD4), align=PP_ALIGN.CENTER, spacing=1.1)
    footer(s)
    notes(s,
          "Quality is the difference between a platform people trust and one they ignore. Walk the four "
          "consequences of poor quality: COSTLY MISTAKES (wrong data → wrong decisions → financial loss), "
          "BROKEN TRUST (one bad dashboard and stakeholders stop believing all of them), WASTED TIME (endless "
          "reconciliation), and COMPLIANCE RISK (inaccurate or unauditable data). The principle on the right — "
          "'garbage in, garbage out' — is the heart of it: processing can clean and reshape data, but it can't "
          "invent correctness that wasn't there. That's why quality checks belong throughout the pipeline, "
          "which the next slide breaks into measurable dimensions.")


def quality_dimensions(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "The Six Dimensions of Data Quality")
    progress(s, 4)
    nodes = [
        dict(title="Accuracy", icon="target", color=LAVA, desc="Reflects reality"),
        dict(title="Completeness", icon="check", color=BLUE, desc="No missing values"),
        dict(title="Consistency", icon="transaction", color=TEAL, desc="Agrees across systems"),
        dict(title="Timeliness", icon="clock", color=PURPLE, desc="Fresh & up to date"),
        dict(title="Validity", icon="shield", color=GREEN, desc="Conforms to rules"),
        dict(title="Uniqueness", icon="variety", color=YELLOW, desc="No duplicates"),
    ]
    radial_infographic(s, SW / 2, 4.4, 2.05, "DATA\nQUALITY", nodes,
                       center_color=LAVA, center_d=1.75, node_d=1.15, label_size=12)
    footer(s)
    notes(s,
          "Data quality isn't vague — it's measurable across six dimensions, shown here as a wheel. ACCURACY: "
          "does the data reflect reality? COMPLETENESS: are values present, not missing? CONSISTENCY: does it "
          "agree across systems and over time? TIMELINESS: is it fresh enough for its use? VALIDITY: does it "
          "conform to the rules, formats and ranges it should? UNIQUENESS: are there no unintended duplicates? "
          "Great pipelines measure and enforce these continuously — and Delta's schema enforcement, "
          "constraints and expectations (in Delta Live Tables / Lakeflow) automate many of them. Give the room "
          "a concrete example for each, e.g. a phone number that's valid in format but inaccurate in value.")


def storage_comparison(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "Warehouse vs. Lake vs. Lakehouse")
    progress(s, 4)
    cols = [
        dict(name="Data Warehouse", color=BLUE, icon="table",
             rows=["Structured data only", "Fast SQL & BI", "Rigid & costly to scale", "No ML / raw data"]),
        dict(name="Data Lake", color=TEAL, icon="cloud",
             rows=["Any data, cheap & scalable", "Great for ML", "No reliability / ACID", "Becomes a swamp"]),
        dict(name="Lakehouse", color=LAVA, icon="delta",
             rows=["All data, one platform", "ACID on cheap storage", "BI + ML together", "Open & governed"]),
    ]
    cw = (CW - 2 * 0.4) / 3
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.4)
        winner = (i == 2)
        card(s, x, 2.05, cw, 4.0, fill=(NAVY if winner else WHITE), line=(None if winner else LINE), shadow=True)
        band(s, x, 2.05, cw, 0.72, c["color"])
        icon(s, c["icon"], x + 0.55, 2.41, 0.42, WHITE)
        text(s, x + 1.0, 2.05, cw - 1.05, 0.72, c["name"], size=13.5, color=WHITE,
             bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.92)
        mk = LAVA_LT if winner else c["color"]
        tc = RGBColor(0xD6, 0xDE, 0xE1) if winner else GRAY
        bullets(s, x + 0.32, 3.0, cw - 0.6, 2.8, c["rows"], size=11, marker=mk, color=tc, gap=10)
        if winner:
            chip(s, x + 0.3, 5.55, cw - 0.6, 0.42, "BEST OF BOTH", LAVA, WHITE, size=10, radius=0.4)
    footer(s)
    notes(s,
          "Bring together the storage evolution the series has touched on. The DATA WAREHOUSE handles only "
          "structured data with fast SQL, but it's rigid, expensive to scale, and can't do ML or raw data. "
          "The DATA LAKE stores any data cheaply and is great for ML, but lacks reliability and governance and "
          "tends to become a swamp. The LAKEHOUSE (highlighted) combines both — all data on one open platform, "
          "ACID reliability on cheap storage, BI and ML together, fully governed. This is the architecture "
          "Databricks pioneered, and it's why we don't have to choose between a warehouse and a lake anymore. "
          "It sets up the medallion and end-to-end slides that close the day.")


def medallion(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "Medallion Architecture")
    progress(s, 4)
    cols = [
        dict(name="Bronze", color=BRONZE, icon="database",
             rows=["Raw, as-ingested", "Full history", "Source of truth"]),
        dict(name="Silver", color=SILVER, icon="filter",
             rows=["Cleansed & conformed", "De-duplicated, joined", "Trusted & queryable"]),
        dict(name="Gold", color=GOLD, icon="diamond",
             rows=["Business aggregates", "KPIs & ML features", "Consumption-ready"]),
    ]
    cw = (CW - 2 * 0.5) / 3
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.5)
        card(s, x, 2.2, cw, 2.95, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.2, cw, 0.8, c["color"])
        icon(s, c["icon"], x + 0.6, 2.6, 0.44, WHITE)
        text(s, x + 1.05, 2.2, cw - 1.1, 0.8, c["name"], size=18, color=WHITE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + 0.32, 3.2, cw - 0.6, 1.8, c["rows"], size=11.5, marker=c["color"], gap=9)
        if i < 2:
            arrow_h(s, x + cw + 0.08, 3.6, 0.4, LAVA, h=0.26)
    band(s, MX, 5.45, CW, 0.5, OAT_2)
    text(s, MX + 0.3, 5.45, 4.6, 0.5, "Data quality rises left to right", size=11, color=NAVY,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    gb = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + 4.9, 5.54, 6.3, 0.28, fill=None, line=None, adj=0.5)
    grad(gb, BRONZE, GOLD, angle=0)
    shape(s, MSO_SHAPE.RIGHT_ARROW, 11.25, 5.52, 0.66, 0.3, fill=GOLD, line=None)
    text(s, MX, 6.15, CW, 0.4,
         "The medallion is how transformation, quality and storage come together in practice.",
         size=11.5, color=GRAY_2, italic=True, align=PP_ALIGN.CENTER)
    footer(s)
    notes(s,
          "The medallion architecture is the practical home for everything in this deck — it operationalizes "
          "the transformation arc, the quality dimensions and the storage layers all at once. BRONZE holds raw "
          "ingested data (full history, source of truth). SILVER is cleansed, de-duplicated, conformed and "
          "joined — trusted and queryable. GOLD is curated business aggregates, KPIs and ML features — "
          "consumption-ready. Quality rises at each hop. Because (on Databricks) every layer is a Delta table, "
          "each transformation is reliable and reversible. This pattern has appeared every day of the series; "
          "by now the audience should see it as the default shape of a data platform.")


def databricks_flow(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "End-to-End Databricks Processing")
    progress(s, 4)
    stages = [
        dict(t="Sources", sub="DBs · APIs\nIoT · files", icon="globe", col=NAVY_2),
        dict(t="Auto Loader", sub="incremental\ningestion", icon="funnel", col=TEAL),
        dict(t="Delta Lake", sub="reliable\nstorage", icon="delta", col=GREEN),
        dict(t="Spark", sub="distributed\nprocessing", icon="bolt", col=LAVA),
        dict(t="Lakehouse", sub="medallion\ntables", icon="warehouse", col=BLUE),
        dict(t="BI / AI", sub="dashboards\n& models", icon="chart", col=PURPLE),
    ]
    journey(s, stages, y=2.6, h=2.5, title_size=11, sub_size=8.4)
    gov = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.3, CW, 0.72, fill=NAVY, line=None, adj=0.2)
    icon(s, "lock", MX + 0.48, 5.66, 0.32, LAVA_LT)
    text(s, MX + 0.9, 5.3, CW - 1.1, 0.72,
         [[{"t": "Unity Catalog", "color": LAVA_LT, "size": 12, "bold": True},
           {"t": "  governs every stage — and Workflows orchestrate the whole pipeline, end to end.",
            "color": WHITE, "size": 11.5}]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, 6.18, CW, 0.4,
         "Every concept from today — ingestion, processing, transformation, storage — in one Databricks flow.",
         size=11.5, color=GRAY_2, italic=True, align=PP_ALIGN.CENTER)
    footer(s)
    notes(s,
          "The capstone — every concept from the day realized on Databricks. SOURCES are ingested incrementally "
          "by AUTO LOADER into DELTA LAKE (reliable storage), processed by SPARK (distributed processing), "
          "organized as a medallion LAKEHOUSE, and served to BI and AI. Underneath, Unity Catalog governs "
          "every stage and Workflows orchestrate the pipeline. Walk it slowly and explicitly map each box to a "
          "concept covered today: ingestion (batch/streaming), processing (distributed Spark), transformation "
          "(medallion), storage (lakehouse), quality and governance. This single slide is the answer to 'how "
          "does Databricks do data processing?'")


def best_practices(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "Data Processing Best Practices")
    progress(s, 4)
    checklist(s, MX + 0.2, 2.2, CW - 0.4, [
        "Prefer ELT — load raw, then transform in the lakehouse",
        "Structure pipelines with Bronze → Silver → Gold",
        "Build quality checks into every stage, not just the end",
        "Choose batch vs streaming per use case — not by default",
        "Make pipelines idempotent so re-runs are safe",
        "Partition data thoughtfully; avoid tiny-file sprawl",
        "Orchestrate with dependencies, retries & monitoring",
        "Govern and document every dataset from day one",
    ], cols=2, row_h=0.92, size=12.5)
    footer(s)
    notes(s,
          "Distil the day into actionable habits. Prefer ELT and the medallion structure. Embed quality checks "
          "throughout the pipeline rather than discovering problems at the end. Choose batch or streaming "
          "deliberately based on the use case's latency needs. Make pipelines IDEMPOTENT — running them twice "
          "produces the same result — which is essential for safe retries (MERGE and Delta make this easy). "
          "Partition thoughtfully to avoid the small-file problem. Orchestrate properly with dependencies, "
          "retries and monitoring. And govern and document datasets from the start. These eight habits "
          "separate fragile, hand-run scripts from robust production data platforms.")


def use_cases(prs):
    s = slide(prs)
    header(s, "Quality & Storage", "Data Processing in the Real World")
    progress(s, 4)
    items = [
        dict(icon="chart", color=BLUE, title="Retail", desc="Demand forecasting & inventory from sales + supply data."),
        dict(icon="lock", color=GREEN, title="Banking", desc="Real-time fraud scoring on streaming transactions."),
        dict(icon="stream", color=LAVA, title="Telecom", desc="Network telemetry processed for live optimization."),
        dict(icon="ml", color=TEAL, title="Healthcare", desc="Unifying records & imaging for analytics and AI."),
        dict(icon="iot", color=PURPLE, title="Manufacturing", desc="IoT sensor pipelines for predictive maintenance."),
        dict(icon="target", color=YELLOW, title="E-Commerce", desc="Clickstream + orders for live recommendations."),
    ]
    grid_cards(s, items, cols=3, x=MX, y=2.0, w=CW, h=4.5, icon_d=0.72,
               title_size=14, desc_size=10.8)
    footer(s)
    notes(s,
          "Close the concepts with industry proof points, each combining the day's ideas. Retail blends batch "
          "sales with supply data for forecasting. Banking runs streaming fraud detection. Telecom processes "
          "massive network telemetry in real time. Healthcare unifies structured records with unstructured "
          "imaging. Manufacturing builds IoT pipelines for predictive maintenance. E-commerce joins "
          "clickstream and orders for live recommendations. Pick the two closest to your audience and trace "
          "the full journey — sources, ingestion mode, transformations, storage, analytics — so they see the "
          "concepts working together. Every one of these runs on the same processing fundamentals.")


def takeaways(prs):
    s = slide(prs)
    header(s, "Wrap-Up", "Key Takeaways")
    cols = [
        dict(no="1", color=LAVA, title="The Journey", points=[
            "Generate → … → Act", "Raw data becomes value"]),
        dict(no="2", color=TEAL, title="Ingest & Move", points=[
            "Batch vs streaming", "ELT > ETL in the cloud"]),
        dict(no="3", color=BLUE, title="Process at Scale", points=[
            "Scale out, not up", "Spark partitions the work"]),
        dict(no="4", color=PURPLE, title="Quality & Store", points=[
            "Six quality dimensions", "Medallion on the Lakehouse"]),
    ]
    cw = (CW - 3 * 0.3) / 4
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.3)
        card(s, x, 2.0, cw, 2.55, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.0, cw, 0.12, c["color"])
        number_badge(s, x + cw / 2, 2.62, 0.62, c["no"], fill=c["color"], size=17)
        text(s, x + 0.1, 3.05, cw - 0.2, 0.4, c["title"], size=12.5, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
        bullets(s, x + 0.26, 3.5, cw - 0.5, 1.0, c["points"], size=10.3, marker=c["color"], gap=7)
    hw = (CW - 0.3) / 2
    takeaway(s, MX, 4.95, hw, "Data processing turns raw signals into trusted, valuable insight.",
             color=LAVA, icon_kind="target")
    takeaway(s, MX + hw + 0.3, 4.95, hw,
             "Scale out with distributed processing — Spark partitions the work.",
             color=BLUE, icon_kind="cluster")
    takeaway(s, MX, 5.85, hw, "Pipelines automate ingest → process → store → serve, reliably.",
             color=TEAL, icon_kind="flow")
    takeaway(s, MX + hw + 0.3, 5.85, hw,
             "Quality + the medallion on the Lakehouse make data production-ready.",
             color=PURPLE, icon_kind="check")
    footer(s)
    notes(s,
          "Recap the arc. THE JOURNEY: data flows generate→…→act, turning raw signals into value. INGEST & "
          "MOVE: choose batch vs streaming, and prefer ELT over ETL in the cloud. PROCESS AT SCALE: scale out, "
          "not up — Spark partitions work across a cluster. QUALITY & STORE: measure the six quality "
          "dimensions and land data in a medallion on the Lakehouse. Quiz the room on the four banner "
          "statements. If they can sketch the data journey and explain why we scale out and prefer ELT, the "
          "day succeeded. Transition to Q&A.")


def closing(prs):
    s = slide(prs, bg=NAVY)
    dark_bg(s)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 2.2, 0.09, 1.5, fill=LAVA, line=None, adj=0.5)
    text(s, MX + 0.3, 2.05, 10, 1.1, "Questions?", size=52, color=WHITE, bold=True)
    text(s, MX + 0.32, 3.25, 10.6, 0.6,
         "You can now trace data end to end — from source, through ingestion and processing, to a governed, "
         "analytics-ready Lakehouse.", size=15, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.2)
    text(s, MX + 0.32, 4.25, 8, 0.3, "GO DEEPER", size=11.5, color=LAVA_LT, bold=True, font=FONT_SB)
    res = ["docs.databricks.com", "Databricks Academy", "Free Edition — hands-on", "Lakeflow & DLT docs"]
    cx = MX + 0.32
    for r in res:
        w = 0.32 + len(r) * 0.105
        chip(s, cx, 4.62, w, 0.52, r, NAVY_2, WHITE, size=11, radius=0.5, line=LINE_DK)
        cx += w + 0.22
    recap = [("Journey", "target", LAVA), ("Ingestion", "funnel", TEAL), ("Transform", "gear", BLUE),
             ("Spark", "bolt", YELLOW), ("Quality", "shield", GREEN), ("Lakehouse", "warehouse", PURPLE)]
    cx = MX + 0.32
    for name, ic, col in recap:
        icon_tile(s, cx + 0.4, 5.85, 0.78, ic, col, shadow=True)
        text(s, cx - 0.1, 6.34, 1.0, 0.4, name, size=9.3, color=RGBColor(0xC4, 0xD0, 0xD4),
             align=PP_ALIGN.CENTER, spacing=0.9)
        cx += 1.15
    text(s, 8.7, 6.55, CW + MX - 8.7, 0.3, "Next: hands-on pipeline lab",
         size=12, color=GRAY_LT, italic=True, align=PP_ALIGN.RIGHT)
    notes(s,
          "Close the day and the conceptual track. Recap the six icons — the journey, ingestion, "
          "transformation, Spark, quality and the Lakehouse. Point to resources: the docs, Databricks Academy, "
          "the free edition for hands-on practice, and the Lakeflow / Delta Live Tables docs for building "
          "declarative pipelines. Preview the next session — a hands-on lab where they build an end-to-end "
          "medallion pipeline with ingestion, transformation and orchestration. Open the floor for questions.")


# ===========================================================================
def main():
    prs = new_deck()
    prs.core_properties.title = "Databricks Day 4 — Data Processing Concepts"
    prs.core_properties.subject = "How data flows from ingestion to analytics"

    cover(prs)
    why_matters(prs)
    data_journey(prs)
    s_divider(prs, "01", "Part One", "Data Sources & Types",
              ["Where enterprise data comes from", "Structured, semi & unstructured data"],
              "globe", color=LAVA)
    sources(prs)
    data_types(prs)
    s_divider(prs, "02", "Part Two", "Data Ingestion",
              ["What ingestion is · batch & streaming", "Choosing the right mode"],
              "funnel", color=LAVA)
    what_ingestion(prs)
    batch_ingestion(prs)
    streaming_ingestion(prs)
    batch_vs_streaming(prs)
    s_divider(prs, "03", "Part Three", "Transformation · ETL & ELT",
              ["Cleaning, standardizing & enriching data", "ETL vs ELT in the cloud"],
              "gear", color=LAVA)
    what_transformation(prs)
    common_transformations(prs)
    etl(prs)
    elt(prs)
    etl_vs_elt(prs)
    s_divider(prs, "04", "Part Four", "Distributed Processing & Pipelines",
              ["Why scale out · the Spark model", "Pipelines & orchestration"],
              "cluster", color=LAVA)
    why_distributed(prs)
    spark_model(prs)
    how_spark_processes(prs)
    what_pipeline(prs)
    pipeline_arch(prs)
    orchestration(prs)
    s_divider(prs, "05", "Part Five", "Quality, Storage & Databricks",
              ["Data quality & its six dimensions", "Storage, medallion & the end-to-end flow"],
              "shield", color=LAVA)
    why_quality(prs)
    quality_dimensions(prs)
    storage_comparison(prs)
    medallion(prs)
    databricks_flow(prs)
    best_practices(prs)
    use_cases(prs)
    takeaways(prs)
    closing(prs)

    prs.save(OUT)
    print(f"Saved {OUT} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    main()
