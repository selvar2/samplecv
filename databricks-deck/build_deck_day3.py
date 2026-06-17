#!/usr/bin/env python3
"""
build_deck_day3.py
------------------
DAY 3 of the Databricks training series — a premium, infographic-rich deck
focused entirely on DELTA LAKE. Same design system as Days 1-2.

Run:  python3 build_deck_day3.py
Out:  Databricks-Day3-Delta-Lake.pptx
"""

from math import ceil
from deck_kit import *
from build_deck import (flow_steps, grid_cards, pipeline, timeline,
                        vs_table, dark_bg, takeaway, s_divider)
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "Databricks-Day3-Delta-Lake.pptx"
PARTS = ["Problem", "Fundamentals", "Internals", "Platform"]


def progress(s, active):
    x = SW - MX - 5.0
    for i, lab in enumerate(PARTS):
        on = (i == active)
        cxp = x + i * 1.25
        circle(s, cxp, 1.46, 0.12 if on else 0.085, LAVA if on else OAT_3)
        text(s, cxp - 0.6, 1.55, 1.2, 0.22, lab, size=7.5,
             color=(NAVY if on else GRAY_LT), align=PP_ALIGN.CENTER, bold=on, font=FONT_SB)


def note_band(s, y, runs, h=1.0, fill=OAT_2):
    card(s, MX, y, CW, h, fill=fill, line=LINE, shadow=False)
    text(s, MX + 0.38, y, CW - 0.76, h, runs, anchor=MSO_ANCHOR.MIDDLE, spacing=1.25)


# ===========================================================================
#  TITLE & AGENDA
# ===========================================================================
def cover(prs):
    s = slide(prs, bg=NAVY)
    dark_bg(s)
    # large Delta triangle motif on the right
    big = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 9.4, 2.1, 3.0, 2.7, fill=None, line=LAVA, line_w=2.5)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 10.15, 3.05, 1.5, 1.4, fill=LAVA, line=None)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 1.45, 0.09, 2.0, fill=LAVA, line=None, adj=0.5)
    text(s, MX + 0.30, 1.40, 9.5, 0.3, "DATABRICKS ENABLEMENT  ·  DAY 3",
         size=12.5, color=LAVA_LT, bold=True, font=FONT_SB)
    text(s, MX + 0.27, 1.82, 9.0, 1.2, "Delta Lake", size=58, color=WHITE, bold=True)
    text(s, MX + 0.29, 2.96, 9.0, 0.8, "Reliability for the Data Lake",
         size=27, color=LAVA, bold=True, font=FONT_LT)
    text(s, MX + 0.30, 3.86, 8.4, 0.7,
         "The open storage layer that brings ACID transactions, time travel and governed "
         "reliability to cheap cloud storage — the foundation of the Lakehouse.",
         size=14, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.25)
    chips = [("01", "The Problem"), ("02", "Fundamentals"),
             ("03", "How It Works"), ("04", "Platform & Value")]
    cx = MX + 0.30
    for no, label in chips:
        w = 2.66
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx, 4.85, w, 0.6, fill=None,
              line=LINE_DK, line_w=1.3, adj=0.5)
        text(s, cx + 0.22, 4.85, w - 0.3, 0.6,
             [[{"t": no + "  ", "color": LAVA, "size": 12, "bold": True},
               {"t": label, "color": WHITE, "size": 11, "bold": True}]],
             anchor=MSO_ANCHOR.MIDDLE)
        cx += w + 0.2
    band(s, MX + 0.30, 6.05, 6.5, 0.02, LINE_DK)
    text(s, MX + 0.30, 6.22, 11, 0.4,
         "Audience:  Data Engineers · Architects · Big Data Developers · Cloud & Analytics Teams",
         size=11.5, color=GRAY_LT)
    text(s, MX + 0.30, 6.62, 11, 0.4,
         "Open source · multi-cloud · the storage foundation beneath every Databricks workload",
         size=11.5, color=GRAY_LT, italic=True)
    notes(s,
          "Welcome to Day 3 — a deep dive on Delta Lake, the single most important storage technology in the "
          "Databricks platform. Frame the arc: Part 1, why traditional data lakes fail and what that costs the "
          "business; Part 2, what Delta Lake is and how it's architected; Part 3, how it works internally — "
          "ACID, the transaction log, time travel, schema enforcement and evolution; Part 4, how it powers the "
          "medallion architecture, the Lakehouse, performance, governance and real pipelines. By the end, "
          "attendees should understand not just what Delta Lake does, but WHY each feature exists and how it "
          "makes a data lake trustworthy.")


def agenda(prs):
    s = slide(prs)
    header(s, "Day 3 Roadmap", "What You'll Learn")
    secs = [
        dict(no="01", color=LAVA, icon="warning", title="The Problem",
             items=["Why data lakes fail", "The cost of bad data quality"],
             outcome="Explain why lakes need a reliability layer."),
        dict(no="02", color=BLUE, icon="delta", title="Fundamentals",
             items=["What is Delta Lake", "Architecture & components"],
             outcome="Describe what Delta Lake is and how it's built."),
        dict(no="03", color=GREEN, icon="transaction", title="How It Works",
             items=["ACID & the transaction log", "Time travel & schema control"],
             outcome="Explain the internals that ensure reliability."),
        dict(no="04", color=PURPLE, icon="gauge", title="Platform & Value",
             items=["Medallion, Lakehouse & performance", "Governance, use cases & best practices"],
             outcome="Apply Delta Lake across an end-to-end platform."),
    ]
    cw = (CW - 3 * 0.3) / 4
    for i, sec in enumerate(secs):
        x = MX + i * (cw + 0.3)
        card(s, x, 2.05, cw, 4.5, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.05, cw, 0.14, sec["color"])
        icon_tile(s, x + cw / 2, 2.95, 0.95, sec["icon"], sec["color"], shadow=True)
        text(s, x + 0.1, 3.55, cw - 0.2, 0.34, "PART " + sec["no"], size=10.5,
             color=sec["color"], bold=True, font=FONT_SB, align=PP_ALIGN.CENTER)
        text(s, x + 0.1, 3.85, cw - 0.2, 0.5, sec["title"], size=15, color=NAVY,
             bold=True, align=PP_ALIGN.CENTER, spacing=0.95)
        band(s, x + 0.35, 4.45, cw - 0.7, 0.014, OAT_3)
        bullets(s, x + 0.3, 4.6, cw - 0.55, 1.2, [(it, 0) for it in sec["items"]],
                size=10.8, marker=sec["color"], gap=7, spacing=1.05)
        text(s, x + 0.3, 5.95, cw - 0.55, 0.5,
             [[{"t": "→ ", "color": sec["color"], "bold": True, "size": 10},
               {"t": sec["outcome"], "color": GRAY_2, "size": 9.5, "italic": True}]], spacing=1.05)
    footer(s)
    notes(s,
          "Lay out the four parts and the learning outcome of each. Stress the narrative logic: we start with "
          "the pain (why lakes fail), then the solution (what Delta is), then the mechanism (how it works), "
          "then the application (how it powers a real platform). Tell learners that Day 1 and Day 2 introduced "
          "Delta Lake at a high level; today we go deep enough that they can reason about reliability, "
          "versioning and performance themselves.")


# ===========================================================================
#  PART 1 — THE PROBLEM
# ===========================================================================
def why_lakes_fail(prs):
    s = slide(prs)
    header(s, "The Problem", "Why Traditional Data Lakes Fail")
    progress(s, 0)
    # broken-lake banner
    card(s, MX, 2.0, CW, 1.05, fill=NAVY, line=None, shadow=True)
    icon(s, "lake", MX + 0.6, 2.52, 0.5, RGBColor(0x6B, 0x77, 0x7C))
    icon(s, "warning", MX + 1.35, 2.52, 0.42, LAVA)
    text(s, MX + 1.95, 2.0, CW - 2.3, 1.05,
         [[{"t": "A raw data lake is just files on object storage — with no transactions, no schema control "
                 "and no guarantees. ", "color": WHITE, "size": 13},
           {"t": "Over time it degrades into a “data swamp.”", "color": LAVA_LT, "size": 13, "bold": True}]],
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    problems = [
        dict(icon="warning", color=LAVA, title="Dirty Data",
             desc="Bad, null & malformed records land unchecked."),
        dict(icon="bolt", color=LAVA_DK, title="Corruption",
             desc="Failed/partial writes leave broken files."),
        dict(icon="transaction", color=RGBColor(0x9C, 0x2A, 0x1E), title="No Transactions",
             desc="Concurrent writes clash — no ACID safety."),
        dict(icon="variety", color=YELLOW, title="Duplicates",
             desc="Re-runs and retries create duplicate rows."),
        dict(icon="branch", color=RGBColor(0xB0, 0x6A, 0x3B), title="Schema Drift",
             desc="Inconsistent, changing schemas break reads."),
    ]
    grid_cards(s, problems, cols=5, x=MX, y=3.35, w=CW, h=3.1, gx=0.24,
               icon_d=0.74, title_size=12.5, desc_size=9.5)
    footer(s)
    notes(s,
          "Set up the entire course problem. A 'data lake' in its raw form is simply a pile of files (often "
          "Parquet) on cheap object storage. That gives scale and low cost, but none of the guarantees a "
          "database provides. Walk the five failure modes: DIRTY DATA (no validation, so nulls and malformed "
          "rows accumulate), CORRUPTION (a job that fails mid-write leaves partial files), NO TRANSACTIONS "
          "(two writers at once can clobber each other — there's no ACID), DUPLICATES (a retried or re-run job "
          "appends the same data twice), and SCHEMA DRIFT (the shape of the data changes silently and breaks "
          "downstream reads). The vivid term is 'data swamp.' Delta Lake exists to fix every one of these.")


def business_impact(prs):
    s = slide(prs)
    header(s, "The Problem", "The Business Cost of Bad Data")
    progress(s, 0)
    items = [
        dict(icon="coins", color=LAVA, title="Revenue Loss",
             desc="Poor data quality costs organizations millions in lost revenue and rework each year."),
        dict(icon="chart", color=BLUE, title="Reporting Errors",
             desc="Dashboards built on unreliable data drive wrong, costly business decisions."),
        dict(icon="clock", color=YELLOW, title="Delayed Analytics",
             desc="Engineers spend time fire-fighting data issues instead of delivering insight."),
        dict(icon="lock", color=PURPLE, title="Governance & Risk",
             desc="No audit trail or lineage means compliance gaps and regulatory exposure."),
    ]
    grid_cards(s, items, cols=2, x=MX, y=2.0, w=7.5, h=4.45, gx=0.3, gy=0.3,
               icon_d=0.8, title_size=14.5, desc_size=11)
    # stat panel
    card(s, 8.3, 2.0, CW + MX - 8.3, 4.45, fill=NAVY, line=None, shadow=True)
    text(s, 8.65, 2.25, 4.0, 0.3, "WHY IT MATTERS", size=11, color=LAVA_LT, bold=True, font=FONT_SB)
    stats = [("~$12.9M", "average annual cost of poor\ndata quality per organization"),
             ("80%", "of an engineer's time spent\nfinding & fixing data issues"),
             ("1 in 3", "leaders distrust the data\nthey use to make decisions")]
    yy = 2.7
    for big, lab in stats:
        text(s, 8.65, yy, 4.0, 0.5, big, size=27, color=WHITE, bold=True)
        text(s, 8.65, yy + 0.5, 4.0, 0.6, lab.replace("\n", " "), size=10.5,
             color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.05)
        yy += 1.18
    footer(s)
    notes(s,
          "Translate the technical problems into business consequences for the executives in the room. Bad "
          "data isn't an engineering nuisance — it's a financial and risk problem. Industry studies (Gartner) "
          "peg the average cost of poor data quality around $12.9M per year per organization. Engineers "
          "reportedly spend up to 80% of their time wrangling and fixing data rather than creating value. And "
          "when leaders don't trust the numbers, analytics initiatives stall. Present the figures as "
          "widely-cited industry estimates. The point: reliability isn't a 'nice to have' — it directly "
          "protects revenue, speed and trust. That's the value case for Delta Lake.")


# ===========================================================================
#  PART 2 — FUNDAMENTALS
# ===========================================================================
def what_is_delta(prs):
    s = slide(prs)
    header(s, "Fundamentals", "What Is Delta Lake?")
    progress(s, 1)
    # radial wheel (left)
    nodes = [
        dict(title="Reliability", icon="check", color=GREEN),
        dict(title="ACID", icon="transaction", color=LAVA),
        dict(title="Time Travel", icon="clock", color=BLUE),
        dict(title="Streaming", icon="stream", color=TEAL),
        dict(title="Governance", icon="lock", color=PURPLE),
        dict(title="Scalability", icon="cluster", color=YELLOW),
    ]
    radial_infographic(s, 3.85, 4.35, 2.0, "DELTA\nLAKE", nodes,
                       center_color=LAVA, center_d=1.7, node_d=1.05, label_size=11)
    # definition (right)
    card(s, 7.85, 2.05, CW + MX - 7.85, 2.5, fill=NAVY, line=None, shadow=True)
    text(s, 8.2, 2.3, 4.3, 0.3, "DEFINITION", size=11, color=LAVA_LT, bold=True, font=FONT_SB)
    text(s, 8.2, 2.64, 4.3, 1.8,
         [[{"t": "Delta Lake is an ", "color": WHITE, "size": 13.5},
           {"t": "open-source storage layer", "color": LAVA_LT, "size": 13.5, "bold": True},
           {"t": " that brings ", "color": WHITE, "size": 13.5},
           {"t": "ACID transactions and reliability", "color": LAVA_LT, "size": 13.5, "bold": True},
           {"t": " to data lakes — turning files on cloud storage into dependable, governed tables.",
            "color": WHITE, "size": 13.5}]], spacing=1.28)
    bullets(s, 7.9, 4.75, CW + MX - 7.85 - 0.1, 1.8, [
        "100% open format — built on Parquet",
        "Unifies batch & streaming on one table",
        "The storage foundation of the Lakehouse",
    ], size=12, marker=LAVA, gap=9)
    footer(s)
    notes(s,
          "Give the crisp definition: Delta Lake is an open-source storage layer that adds ACID transactions "
          "and reliability on top of data lakes. The radial wheel previews the six capabilities we'll explore "
          "— reliability, ACID, time travel, streaming, governance and scalability — all radiating from one "
          "open table format. Three things to stress: it's OPEN (just Parquet files plus a log, no "
          "proprietary lock-in), it UNIFIES batch and streaming on a single table, and it's the foundation "
          "the entire Lakehouse is built on. Everything in Parts 3 and 4 elaborates one of these spokes.")


def delta_architecture(prs):
    s = slide(prs)
    header(s, "Fundamentals", "Delta Lake Architecture")
    progress(s, 1)
    # left: vertical stack Users -> Databricks -> Delta Lake -> Cloud Storage
    layers = [("Users", "Engineers · Analysts · Scientists · BI", "people", NAVY_2),
              ("Databricks", "Spark + Photon read & write Delta", "bolt", LAVA),
              ("Delta Lake", "Open transactional table layer", "delta", GREEN),
              ("Cloud Storage", "S3 · ADLS · GCS — your account", "cloud", BLUE)]
    for i, (t, d, ic, col) in enumerate(layers):
        y = 2.05 + i * 1.16
        card(s, MX, y, 6.4, 1.0, fill=WHITE, line=LINE, shadow=True)
        band(s, MX, y, 0.1, 1.0, col)
        icon_tile(s, MX + 0.65, y + 0.5, 0.66, ic, col, shadow=False)
        text(s, MX + 1.15, y + 0.13, 5.1, 0.4, t, size=14, color=NAVY, bold=True)
        text(s, MX + 1.15, y + 0.52, 5.1, 0.4, d, size=10.5, color=GRAY_2, spacing=1.0)
        if i < 3:
            shape(s, MSO_SHAPE.DOWN_ARROW, MX + 3.1, y + 1.0, 0.2, 0.18, fill=GRAY_LT, line=None)
    # right: what's inside a Delta table
    card(s, 7.5, 2.05, CW + MX - 7.5, 4.4, fill=NAVY, line=None, shadow=True)
    text(s, 7.85, 2.3, 4.5, 0.3, "INSIDE A DELTA TABLE", size=11, color=LAVA_LT, bold=True, font=FONT_SB)
    inside = [("Transaction Log", "_delta_log/ — ordered JSON commits; the brain of the table", "transaction", LAVA),
              ("Metadata Layer", "Schema, partitions, statistics & file list", "doc", BLUE),
              ("Data Files", "Immutable Parquet files on object storage", "cube", GREEN)]
    for i, (t, d, ic, col) in enumerate(inside):
        y = 2.74 + i * 1.18
        sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.85, y, CW + MX - 7.5 - 0.7, 1.0,
                   fill=NAVY_2, line=col, line_w=1.4, adj=0.1)
        icon_tile(s, 8.3, y + 0.5, 0.6, ic, col, shadow=False)
        text(s, 8.75, y + 0.12, 3.5, 0.36, t, size=12.5, color=WHITE, bold=True)
        text(s, 8.75, y + 0.46, 3.6, 0.5, d, size=9.3, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.0)
    footer(s)
    notes(s,
          "Show how Delta Lake sits in the stack and what a Delta table actually contains. On the left, the "
          "layered flow: USERS work through DATABRICKS (Spark + Photon), which reads and writes DELTA LAKE "
          "tables that physically live in the customer's own CLOUD STORAGE. Delta is the reliability layer "
          "between the compute and the raw files. On the right, open up a single Delta table: it's three "
          "things — the TRANSACTION LOG (_delta_log, the ordered JSON commits that are the brain of the "
          "table), the METADATA (schema, partitions, statistics, the list of valid files), and the DATA FILES "
          "(immutable Parquet). The log is what we'll spend Part 3 on — it's the source of every guarantee.")


def delta_components(prs):
    s = slide(prs)
    header(s, "Fundamentals", "The Five Core Components")
    progress(s, 1)
    nodes = [
        dict(title="Delta Table", icon="table", color=LAVA, desc="The reliable unit"),
        dict(title="Transaction Log", icon="transaction", color=BLUE, desc="Ordered commits"),
        dict(title="Metadata", icon="doc", color=GREEN, desc="Schema & stats"),
        dict(title="Storage Layer", icon="cube", color=TEAL, desc="Parquet files"),
        dict(title="Optimization", icon="gauge", color=PURPLE, desc="OPTIMIZE · Z-ORDER"),
    ]
    radial_infographic(s, SW / 2, 4.35, 2.05, "DELTA\nLAKE", nodes,
                       center_color=NAVY, center_d=1.75, node_d=1.15, label_size=12)
    footer(s)
    notes(s,
          "Summarize the moving parts as a wheel. The DELTA TABLE is the user-facing unit — what feels like a "
          "single reliable table. Underneath, the TRANSACTION LOG records every change as ordered commits. "
          "The METADATA layer tracks schema, partitions and column statistics (used to skip files for speed). "
          "The STORAGE LAYER is the immutable Parquet data files on object storage. And the OPTIMIZATION "
          "ENGINE provides operations like OPTIMIZE (compaction) and Z-ORDER (data clustering) to keep reads "
          "fast. These five pieces working together are what make a folder of files behave like a database. "
          "We'll see the log and optimization in detail shortly.")


# ===========================================================================
#  PART 3 — HOW IT WORKS
# ===========================================================================
def acid(prs):
    s = slide(prs)
    header(s, "How It Works", "ACID Transactions Explained")
    progress(s, 2)
    cards_def = [
        dict(L="A", title="Atomicity", icon="transaction", color=LAVA,
             desc="All-or-nothing — a write fully succeeds or leaves the table untouched."),
        dict(L="C", title="Consistency", icon="check", color=BLUE,
             desc="Every transaction moves the table from one valid state to another."),
        dict(L="I", title="Isolation", icon="lock", color=GREEN,
             desc="Concurrent reads & writes don't interfere — no dirty reads."),
        dict(L="D", title="Durability", icon="cube", color=PURPLE,
             desc="Once committed, data is permanent — survives any failure."),
    ]
    cw = (CW - 3 * 0.3) / 4
    for i, c in enumerate(cards_def):
        x = MX + i * (cw + 0.3)
        card(s, x, 2.1, cw, 3.7, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.1, cw, 0.12, c["color"])
        # big letter
        circle(s, x + cw / 2, 2.95, 1.0, c["color"], shadow=True)
        text(s, x + cw / 2 - 0.5, 2.5, 1.0, 0.9, c["L"], size=40, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.12, 3.62, cw - 0.24, 0.4, c["title"], size=14, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
        text(s, x + 0.22, 4.08, cw - 0.44, 1.6, c["desc"], size=10.5, color=GRAY_2,
             align=PP_ALIGN.CENTER, spacing=1.12)
        icon(s, c["icon"], x + cw / 2, 5.45, 0.4, c["color"])
    note_band(s, 6.1,
              [[{"t": "ACID is what separates a database from a pile of files.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Delta Lake brings all four guarantees to cheap object storage — no warehouse required.",
                 "color": GRAY, "size": 12.5}]], h=0.85)
    footer(s)
    notes(s,
          "ACID is the bedrock concept — define each with a relatable example. ATOMICITY: a write is "
          "all-or-nothing; if a job loading a million rows fails at row 600k, the table is left exactly as it "
          "was — no half-loaded mess. CONSISTENCY: every committed transaction leaves the table in a valid "
          "state that respects its schema and constraints. ISOLATION: while one job is writing, readers still "
          "see the last good version — they never see a dirty, in-progress state, and concurrent writers are "
          "serialized safely. DURABILITY: once a commit succeeds, it's permanent and survives crashes. The "
          "punchline: these four properties are exactly what traditional data lakes lacked, and what Delta "
          "delivers directly on object storage.")


def transaction_log(prs):
    s = slide(prs)
    header(s, "How It Works", "The Transaction Log  (_delta_log)")
    progress(s, 2)
    steps = [
        dict(title="Write Request", icon="doc", color=NAVY_2, badge=1,
             desc="A job proposes changes to the table."),
        dict(title="Log Update", icon="transaction", color=BLUE, badge=2,
             desc="A new commit is prepared in _delta_log."),
        dict(title="Atomic Commit", icon="check", color=GREEN, badge=3,
             desc="The commit succeeds all-at-once, or not at all."),
        dict(title="Read Consistency", icon="people", color=LAVA, badge=4,
             desc="Readers replay the log → a consistent snapshot."),
    ]
    flow_steps(s, steps, y_top=2.3, card_h=2.35, icon_d=0.96, arrow=True, desc_dy=1.74)
    # log versions strip
    log = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.05, CW, 1.05, fill=NAVY, line=None, adj=0.08)
    text(s, MX + 0.35, 5.16, 5.0, 0.3, "_delta_log/   ordered, immutable JSON commits", size=10.5,
         color=LAVA_LT, bold=True, font=FONT_SB)
    versions = ["00000.json", "00001.json", "00002.json", "00003.json", "…"]
    vw = (CW - 0.7 - 4 * 0.18) / 5
    for i, v in enumerate(versions):
        x = MX + 0.35 + i * (vw + 0.18)
        cur = (i == 3)
        sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.52, vw, 0.46,
                   fill=(LAVA if cur else NAVY_2), line=(None if cur else LINE_DK), line_w=1, adj=0.16)
        text(s, x + 0.02, 5.52, vw - 0.04, 0.46, v, size=9.5, color=WHITE, bold=cur,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 4:
            arrow_h(s, x + vw + 0.0, 5.75, 0.18, RGBColor(0x7E, 0x91, 0x99), h=0.12)
    footer(s)
    notes(s,
          "The transaction log is the single most important internal concept — everything else derives from "
          "it. Walk the flow: a WRITE REQUEST proposes changes; Delta prepares a new commit (LOG UPDATE) that "
          "records exactly which files are added and removed; the ATOMIC COMMIT writes a new numbered JSON "
          "file into the _delta_log folder in one operation — it either lands completely or not at all; and "
          "any reader achieves READ CONSISTENCY by replaying the ordered log to compute the current set of "
          "valid files, giving a clean snapshot. Show the version strip: 00000.json, 00001.json, … — each "
          "commit is a new immutable version. This is simultaneously the source of ACID, time travel and "
          "audit. Optimistic concurrency control resolves competing writers.")


def reliability_flow(prs):
    s = slide(prs)
    header(s, "How It Works", "How Delta Ensures Reliability")
    progress(s, 2)
    steps = [
        dict(title="Write", icon="doc", color=NAVY_2, badge=1,
             desc="New data files are written to storage."),
        dict(title="Validate", icon="shield", color=YELLOW, badge=2,
             desc="Schema & constraints are enforced."),
        dict(title="Commit", icon="check", color=GREEN, badge=3,
             desc="Atomic commit recorded in the log."),
        dict(title="Version", icon="clock", color=BLUE, badge=4,
             desc="A new table version is created."),
        dict(title="Read", icon="people", color=LAVA, badge=5,
             desc="Consistent snapshot served to all."),
    ]
    flow_steps(s, steps, y_top=2.35, card_h=2.5, icon_d=1.0, arrow=True)
    note_band(s, 5.25,
              [[{"t": "Every write follows the same safe path.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Validation rejects bad data, the atomic commit prevents corruption, and versioning "
                      "makes every change reversible — reliability by design, not by hope.",
                 "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Tie ACID and the log together into the end-to-end reliability path that every write follows. WRITE: "
          "new Parquet files are staged. VALIDATE: the data is checked against the table's schema and "
          "constraints — bad data is rejected before it can pollute the table. COMMIT: an atomic entry is "
          "added to the transaction log. VERSION: that commit becomes a new, numbered version of the table. "
          "READ: every reader gets a consistent snapshot computed from the log. The message for engineers: "
          "reliability isn't something you bolt on with extra code — it's structurally guaranteed by this "
          "pipeline. This is exactly what raw data lakes couldn't promise.")


def time_travel(prs):
    s = slide(prs)
    header(s, "How It Works", "Time Travel & Data Versioning")
    progress(s, 2)
    # version timeline
    vers = [("v0", "Initial load", NAVY_2), ("v1", "+ new rows", BLUE),
            ("v2", "Bad update", LAVA_DK), ("v3", "Rollback to v1", GREEN)]
    n = len(vers); aw = 0.5
    bw = (CW - (n - 1) * aw) / n
    for i, (v, lab, col) in enumerate(vers):
        x = MX + i * (bw + aw)
        card(s, x, 2.35, bw, 1.7, fill=WHITE, line=LINE, shadow=True)
        circle(s, x + bw / 2, 2.95, 0.78, col, shadow=True)
        text(s, x + bw / 2 - 0.5, 2.62, 1.0, 0.7, v, size=22, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.1, 3.55, bw - 0.2, 0.4, lab, size=11.5, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
        if i < n - 1:
            arrow_h(s, x + bw + 0.06, 2.95, aw - 0.12, GRAY_LT, h=0.2)
    # rollback arc note
    rb = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + bw + aw - 0.1, 4.35, bw * 2 + aw + 0.2, 0.6,
               fill=None, line=GREEN, line_w=1.6, adj=0.3)
    text(s, MX + bw + aw - 0.1, 4.35, bw * 2 + aw + 0.2, 0.6,
         "↶  Restore any prior version in one command", size=11.5, color=GREEN_DK, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    note_band(s, 5.3,
              [[{"t": "Every commit is a version.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Query the table “AS OF” a timestamp or version for audits, reproducible ML and "
                      "debugging — or instantly RESTORE after a bad write. Mistakes become undoable.",
                 "color": GRAY, "size": 12.5}]], h=0.95)
    footer(s)
    notes(s,
          "Time travel is one of Delta's most loved features and follows directly from the log. Because every "
          "commit creates a new version (v0, v1, v2, …), you can query the table AS OF any past version or "
          "timestamp. Walk the story: v0 initial load, v1 adds rows, v2 is a bad update that corrupts values, "
          "and v3 simply RESTORES back to v1 — the bad change is undone in one command. Real uses: auditing "
          "(what did this table look like last quarter?), reproducible ML (train on the exact data snapshot), "
          "debugging (compare versions), and instant recovery from mistakes. The reassuring message: with "
          "Delta, a bad write is no longer a disaster — it's reversible.")


def schema_enforcement(prs):
    s = slide(prs)
    header(s, "How It Works", "Schema Enforcement")
    progress(s, 2)
    before_after(s, 2.05, 3.5,
                 "Without Delta", [
                     "Bad / mistyped data written silently",
                     "Schema drifts with every job",
                     "Downstream reads break unexpectedly",
                     "No guardrails — garbage in, garbage out"],
                 "With Delta", [
                     "Writes validated against the schema",
                     "Mismatched data is rejected on write",
                     "Tables stay clean & predictable",
                     "Quality enforced automatically"])
    note_band(s, 5.85,
              [[{"t": "Schema enforcement is the gatekeeper.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Delta refuses writes that don't match the table's schema — preventing the silent "
                      "corruption that turns lakes into swamps.", "color": GRAY, "size": 12.5}]], h=0.85)
    footer(s)
    notes(s,
          "Schema enforcement (also called schema-on-write validation) is Delta's quality gatekeeper. Use the "
          "before/after contrast. WITHOUT Delta, any job can write any shape of data — a column typo, a wrong "
          "data type, an extra field — and it lands silently, drifting the schema and eventually breaking the "
          "dashboards and pipelines that read it. WITH Delta, every write is checked against the table's "
          "declared schema; if it doesn't match, the write is rejected before any damage is done. This single "
          "behavior prevents the most common cause of 'data swamp' decay. Note it's strict by default but "
          "intentionally overridable — which leads to the next slide, schema evolution.")


def schema_evolution(prs):
    s = slide(prs)
    header(s, "How It Works", "Schema Evolution")
    progress(s, 2)
    steps = [
        dict(title="New Column", icon="variety", color=NAVY_2, badge=1,
             desc="Source data arrives with a new, intended field."),
        dict(title="Validation", icon="shield", color=YELLOW, badge=2,
             desc="Delta checks the change is safe & explicit."),
        dict(title="Evolution", icon="branch", color=BLUE, badge=3,
             desc="mergeSchema adds the column intentionally."),
        dict(title="Updated Table", icon="check", color=GREEN, badge=4,
             desc="Schema evolves — no pipeline rewrite needed."),
    ]
    flow_steps(s, steps, y_top=2.3, card_h=2.4, icon_d=0.96, arrow=True, desc_dy=1.74)
    note_band(s, 5.1,
              [[{"t": "Enforcement and evolution work together.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Enforcement blocks ", "color": GRAY, "size": 12.5},
                {"t": "accidental", "color": LAVA, "size": 12.5, "bold": True},
                {"t": " changes; evolution allows ", "color": GRAY, "size": 12.5},
                {"t": "intentional", "color": GREEN_DK, "size": 12.5, "bold": True},
                {"t": " ones — so tables adapt safely as the business changes.",
                 "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Schema evolution is the deliberate counterpart to enforcement. Business data legitimately changes — "
          "a new column appears, a field is added. Without control, that's chaos; with Delta, it's a managed "
          "operation. Walk the flow: a NEW COLUMN arrives in the source; Delta VALIDATES that the change is "
          "safe and explicit; using an explicit option (mergeSchema), the table EVOLVES to include the new "
          "column; and you get an UPDATED TABLE with zero pipeline rewrites. The key teaching point: "
          "enforcement and evolution are two sides of one coin — enforcement blocks ACCIDENTAL drift, "
          "evolution permits INTENTIONAL change. Together they let tables adapt without ever becoming "
          "unreliable.")


def batch_streaming(prs):
    s = slide(prs)
    header(s, "How It Works", "Unified Batch + Streaming")
    progress(s, 2)
    # left sources
    srcs = [("Batch Sources", "Files · DB loads · history", "database", BLUE),
            ("Streaming Sources", "Kafka · events · IoT", "stream", LAVA)]
    for i, (t, d, ic, col) in enumerate(srcs):
        y = 2.35 + i * 1.5
        card(s, MX, y, 3.5, 1.25, fill=WHITE, line=LINE, shadow=True)
        band(s, MX, y, 0.1, 1.25, col)
        icon_tile(s, MX + 0.6, y + 0.62, 0.66, ic, col, shadow=False)
        text(s, MX + 1.1, y + 0.2, 2.3, 0.4, t, size=12.5, color=NAVY, bold=True, spacing=0.95)
        text(s, MX + 1.1, y + 0.66, 2.3, 0.5, d, size=9.5, color=GRAY_2, spacing=1.0)
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(MX + 3.5), Inches(y + 0.62),
                                    Inches(5.55), Inches(4.0))
        ln.line.color.rgb = GRAY_LT; ln.line.width = Pt(1.8)
    # center delta
    circle(s, 6.4, 4.0, 1.7, GREEN, line=WHITE, line_w=3, shadow=True)
    icon(s, "delta", 6.4, 3.7, 0.55, WHITE)
    text(s, 5.55, 4.05, 1.7, 0.6, "Delta Lake\nONE table", size=12.5, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
    # right consumers
    cons = [("BI & Reporting", "chart", BLUE), ("Machine Learning", "ml", PURPLE), ("Analytics", "gauge", TEAL)]
    for i, (t, ic, col) in enumerate(cons):
        y = 2.2 + i * 1.25
        card(s, 9.4, y, CW + MX - 9.4, 1.0, fill=WHITE, line=LINE, shadow=True)
        icon_tile(s, 9.4 + 0.5, y + 0.5, 0.6, ic, col, shadow=False)
        text(s, 9.4 + 0.95, y, CW + MX - 9.4 - 1.1, 1.0, t, size=12, color=NAVY, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.25), Inches(4.0),
                                    Inches(9.4), Inches(y + 0.5))
        ln.line.color.rgb = GRAY_LT; ln.line.width = Pt(1.8)
    footer(s)
    notes(s,
          "A defining Delta capability: the same table serves both batch and streaming, eliminating the "
          "classic 'two pipelines' problem (the old Lambda architecture). On the left, BATCH sources (file "
          "loads, database dumps, history) and STREAMING sources (Kafka, events, IoT) both write into ONE "
          "Delta table. On the right, that single table simultaneously feeds BI, machine learning and "
          "analytics. Because Delta provides ACID and a transaction log, a streaming writer and a batch "
          "reader can safely share the same table at the same time. The benefit: one copy of data, one "
          "codebase, no reconciling separate batch and speed layers — a huge simplification.")


# ===========================================================================
#  PART 4 — PLATFORM & VALUE
# ===========================================================================
def medallion(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Delta Powers the Medallion Architecture")
    progress(s, 3)
    cols = [
        dict(name="Bronze", color=BRONZE, icon="database",
             rows=["Raw data, as ingested", "Append-only history", "Source of truth"]),
        dict(name="Silver", color=SILVER, icon="gear",
             rows=["Cleansed & conformed", "De-duplicated, joined", "Trusted & queryable"]),
        dict(name="Gold", color=GOLD, icon="diamond",
             rows=["Business aggregates", "ML features & KPIs", "Consumption-ready"]),
    ]
    cw = (CW - 2 * 0.5) / 3
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.5)
        card(s, x, 2.2, cw, 2.95, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.2, cw, 0.8, c["color"])
        icon(s, c["icon"], x + 0.6, 2.6, 0.44, WHITE)
        text(s, x + 1.05, 2.2, cw - 1.1, 0.8, c["name"], size=18, color=WHITE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + 0.32, 3.2, cw - 0.6, 1.8, c["rows"], size=11.5, marker=c["color"], gap=9)
        if i < 2:
            arrow_h(s, x + cw + 0.08, 3.6, 0.4, LAVA, h=0.26)
    band(s, MX, 5.45, CW, 0.5, OAT_2)
    text(s, MX + 0.3, 5.45, 4.6, 0.5, "Every layer is a Delta table", size=11, color=NAVY,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    gb = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + 4.9, 5.54, 6.3, 0.28, fill=None, line=None, adj=0.5)
    grad(gb, BRONZE, GOLD, angle=0)
    shape(s, MSO_SHAPE.RIGHT_ARROW, 11.25, 5.52, 0.66, 0.3, fill=GOLD, line=None)
    text(s, MX, 6.15, CW, 0.4,
         "Reliability compounds layer by layer — because Bronze, Silver and Gold are all ACID Delta tables.",
         size=11.5, color=GRAY_2, italic=True, align=PP_ALIGN.CENTER)
    footer(s)
    notes(s,
          "Connect Delta to the architecture pattern attendees saw in Days 1-2 — but now they understand WHY "
          "it works. The medallion refines data through Bronze (raw), Silver (cleansed) and Gold (curated) "
          "layers. The crucial point for this deck: every one of those layers is a Delta table, so ACID, "
          "schema enforcement and time travel apply at every hop. That's what makes the pattern trustworthy — "
          "if Silver were just raw files, a failed cleansing job could corrupt it. Because it's Delta, the "
          "transformation either commits cleanly or not at all, and you can always roll back. Reliability "
          "compounds as data flows toward Gold.")


def lakehouse(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Delta Lake + the Lakehouse")
    progress(s, 3)
    parts = [dict(t="Cloud Storage", d="Cheap, scalable object storage", ic="cloud", col=BLUE),
             dict(t="Delta Lake", d="Reliability & ACID layer", ic="delta", col=GREEN),
             dict(t="Databricks", d="Spark, SQL, ML & governance", ic="bolt", col=LAVA)]
    iw, gap = 2.45, 0.5
    xs = [MX + i * (iw + gap) for i in range(3)]          # 0.62, 3.57, 6.52
    for x, p in zip(xs, parts):
        card(s, x, 2.5, iw, 2.4, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.5, iw, 0.12, p["col"])
        icon_tile(s, x + iw / 2, 3.35, 1.0, p["ic"], p["col"], shadow=True)
        text(s, x + 0.1, 3.95, iw - 0.2, 0.4, p["t"], size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.12, 4.36, iw - 0.24, 0.5, p["d"], size=10, color=GRAY_2,
             align=PP_ALIGN.CENTER, spacing=1.0)
    for xc in (MX + iw, MX + 2 * iw + gap):               # the two "+" gaps
        text(s, xc, 3.3, gap, 0.8, "+", size=32, color=GRAY_LT, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    eqx = MX + 3 * iw + 2 * gap                            # the "=" gap (8.97)
    text(s, eqx, 3.3, gap, 0.8, "=", size=32, color=LAVA, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rx = eqx + gap                                         # result card start (9.47)
    rw = CW + MX - rx
    res = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, rx, 2.5, rw, 2.4, fill=NAVY, line=None, adj=0.06)
    soft_shadow(res)
    icon_tile(s, rx + rw / 2, 3.35, 1.05, "warehouse", LAVA, shadow=True)
    text(s, rx, 3.95, rw, 0.45, "LAKEHOUSE", size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, rx, 4.42, rw, 0.4, "lake economics + warehouse reliability",
         size=9.5, color=LAVA_LT, align=PP_ALIGN.CENTER, italic=True)
    note_band(s, 5.35,
              [[{"t": "Delta Lake is the keystone.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "It's the reliability layer that turns cheap cloud storage into a Lakehouse — without "
                      "Delta, you just have a lake.", "color": GRAY, "size": 12.5}]], h=1.0)
    footer(s)
    notes(s,
          "Show where Delta fits in the bigger picture with a simple equation. Cloud Storage gives you cheap, "
          "infinite capacity but no reliability. Databricks gives you powerful compute and governance. Delta "
          "Lake is the missing middle layer — the reliability and transaction layer — that binds them into a "
          "Lakehouse. The key message: Delta is the KEYSTONE. Without it, cloud storage plus compute is just "
          "a data lake with all the problems from Part 1. With it, you get warehouse-grade reliability on "
          "lake-grade economics. This is why Delta is the default table format on Databricks.")


def performance(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Performance Optimization")
    progress(s, 3)
    items = [
        dict(icon="cube", color=LAVA, title="OPTIMIZE",
             desc="Compacts many small files into fewer large ones for faster scans."),
        dict(icon="sort", color=BLUE, title="Z-ORDER",
             desc="Co-locates related data so queries skip irrelevant files."),
        dict(icon="merge", color=GREEN, title="Compaction",
             desc="Auto-compaction & optimized writes tame the small-file problem."),
        dict(icon="bolt", color=PURPLE, title="Caching",
             desc="Delta cache & Photon keep hot data fast and close to compute."),
    ]
    grid_cards(s, items, cols=2, x=MX, y=2.0, w=7.5, h=4.45, gx=0.3, gy=0.3,
               icon_d=0.78, title_size=14, desc_size=10.8)
    # mini performance gauge panel
    card(s, 8.3, 2.0, CW + MX - 8.3, 4.45, fill=NAVY, line=None, shadow=True)
    text(s, 8.65, 2.25, 4.0, 0.3, "THE PAYOFF", size=11, color=LAVA_LT, bold=True, font=FONT_SB)
    icon_tile(s, 8.3 + (CW + MX - 8.3) / 2, 3.5, 1.5, "gauge", LAVA, shadow=True)
    text(s, 8.5, 4.45, CW + MX - 8.3 - 0.4, 0.5, "Faster queries · lower cost", size=12.5,
         color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    bullets(s, 8.6, 5.0, CW + MX - 8.3 - 0.5, 1.4, [
        "Data skipping via file statistics",
        "Fewer, bigger files = less overhead",
        "Photon-native execution",
    ], size=10.3, marker=LAVA_LT, color=RGBColor(0xC4, 0xD0, 0xD4), gap=7)
    footer(s)
    notes(s,
          "Delta isn't just reliable — it's fast, and these four levers are how. OPTIMIZE compacts the many "
          "small files that streaming and frequent writes create into fewer large ones, which dramatically "
          "speeds up scans (the 'small file problem' is a top cause of slow lakes). Z-ORDER physically "
          "co-locates related data so that, combined with the column statistics in the log, queries can SKIP "
          "files that can't match — reading far less data. Auto-compaction and optimized writes do this "
          "automatically. And the Delta cache plus Photon keep hot data fast. The combined payoff: faster "
          "queries at lower cost, because the engine reads less and works more efficiently.")


def governance(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Governance with Unity Catalog")
    progress(s, 3)
    steps = [
        dict(title="Unity Catalog", icon="lock", color=LAVA, badge=1,
             desc="One catalog governs every Delta table & AI asset."),
        dict(title="Security", icon="shield", color=BLUE, badge=2,
             desc="Fine-grained access — row, column & SQL GRANTs."),
        dict(title="Audit & Lineage", icon="branch", color=GREEN, badge=3,
             desc="Automatic column-level lineage and audit logs."),
        dict(title="Compliance", icon="check", color=PURPLE, badge=4,
             desc="GDPR · HIPAA · SOC 2 — with time-travel evidence."),
    ]
    flow_steps(s, steps, y_top=2.3, card_h=2.4, icon_d=0.96, arrow=True, desc_dy=1.74)
    note_band(s, 5.1,
              [[{"t": "Delta + Unity Catalog = governed reliability.  ", "color": NAVY, "size": 12.5, "bold": True},
                {"t": "Delta makes data trustworthy; Unity Catalog controls who can see it, tracks where it "
                      "came from, and proves compliance — together across every cloud.", "color": GRAY, "size": 12.5}]],
              h=1.0)
    footer(s)
    notes(s,
          "Reliability and governance are complementary: Delta makes data trustworthy, Unity Catalog controls "
          "and proves how it's used. Walk the chain: UNITY CATALOG provides one governance layer over every "
          "Delta table (and ML model and file); SECURITY means fine-grained, SQL-based access control down to "
          "rows and columns, tied to corporate identity; AUDIT & LINEAGE are automatic and column-level, so "
          "you can see exactly where data came from and who touched it; and COMPLIANCE (GDPR, HIPAA, SOC 2) is "
          "achievable — Delta's time travel even provides historical evidence of what data looked like at any "
          "point. The combination is what enterprises need to run regulated workloads on a lake.")


def end_to_end(prs):
    s = slide(prs)
    header(s, "Platform & Value", "End-to-End Data Flow on Delta")
    progress(s, 3)
    stages = [
        dict(t="Sources", sub="Apps · DBs · IoT\nfiles · events", icon="globe", col=NAVY_2),
        dict(t="Ingestion", sub="Auto Loader\nstream + batch", icon="funnel", col=TEAL),
        dict(t="Bronze", sub="Raw Delta\nappend-only", icon="database", col=BRONZE),
        dict(t="Silver", sub="Cleansed &\nconformed", icon="gear", col=SILVER),
        dict(t="Gold", sub="Aggregates &\nfeatures", icon="diamond", col=GOLD),
        dict(t="Dashboard", sub="BI · Genie\nSQL analytics", icon="chart", col=BLUE),
        dict(t="AI / ML", sub="Models &\nGenAI", icon="ml", col=PURPLE),
    ]
    n = len(stages); aw = 0.16
    cw = (CW - (n - 1) * aw) / n
    y, h = 2.6, 2.5
    for i, st in enumerate(stages):
        x = MX + i * (cw + aw)
        card(s, x, y, cw, h, fill=WHITE, line=LINE, shadow=True)
        band(s, x, y, cw, 0.1, st["col"])
        icon_tile(s, x + cw / 2, y + 0.7, 0.8, st["icon"], st["col"], shadow=False)
        text(s, x + 0.04, y + 1.2, cw - 0.08, 0.42, st["t"], size=11, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, spacing=0.9)
        text(s, x + 0.04, y + 1.66, cw - 0.08, 0.8, st["sub"], size=8.2, color=GRAY_2,
             align=PP_ALIGN.CENTER, spacing=1.0)
        number_badge(s, x + 0.26, y + 0.26, 0.34, i + 1, fill=st["col"], size=10)
        if i < n - 1:
            arrow_h(s, x + cw + 0.0, y + h / 2, aw, GRAY_LT, h=0.16)
    gov = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.35, CW, 0.72, fill=NAVY, line=None, adj=0.2)
    icon(s, "delta", MX + 0.48, 5.71, 0.34, GREEN)
    text(s, MX + 0.9, 5.35, CW - 1.1, 0.72,
         [[{"t": "Delta Lake", "color": GREEN, "size": 12, "bold": True},
           {"t": "  is the reliable, governed storage under every stage — one open copy of data, end to end.",
            "color": WHITE, "size": 11.5}]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, 6.2, CW, 0.4,
         "Source to insight on a single, reliable foundation — ingested, refined, served and learned from.",
         size=11.5, color=GRAY_2, italic=True, align=PP_ALIGN.CENTER)
    footer(s)
    notes(s,
          "The capstone pipeline, now told from Delta's perspective. Data flows from SOURCES through INGESTION "
          "(Auto Loader handles both streaming and batch) into the BRONZE → SILVER → GOLD medallion, then out "
          "to DASHBOARDS and AI/ML. The single unifying message of the whole day: every stage reads and writes "
          "DELTA LAKE tables, so the entire pipeline inherits ACID reliability, schema control, time travel "
          "and governance from end to end — one open copy of data, no silos, no swamp. Walk it slowly and "
          "connect each box back to a feature covered earlier today.")


def use_cases(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Delta Lake in the Real World")
    progress(s, 3)
    items = [
        dict(icon="coins", color=BLUE, title="Retail",
             desc="Real-time inventory & personalization on one reliable copy of data."),
        dict(icon="lock", color=GREEN, title="Banking",
             desc="Auditable, ACID-compliant transactions with full time-travel history."),
        dict(icon="ml", color=LAVA, title="Healthcare",
             desc="Reproducible patient & research datasets versioned for compliance."),
        dict(icon="gear", color=TEAL, title="Manufacturing",
             desc="IoT sensor streams + batch merged into trusted Delta tables."),
        dict(icon="stream", color=PURPLE, title="Telecom",
             desc="High-volume event streams unified for real-time network analytics."),
        dict(icon="chart", color=YELLOW, title="E-Commerce",
             desc="Clickstream + orders joined reliably for live recommendations."),
    ]
    grid_cards(s, items, cols=3, x=MX, y=2.0, w=CW, h=4.5, icon_d=0.72,
               title_size=14, desc_size=10.8)
    footer(s)
    notes(s,
          "Show breadth across industries, anchoring each to a Delta capability the audience now understands. "
          "Retail and e-commerce lean on unified batch+streaming for real-time inventory and recommendations. "
          "Banking depends on ACID and time travel for auditable, compliant transaction data. Healthcare "
          "values versioning for reproducible research and regulatory evidence. Manufacturing and telecom "
          "rely on merging massive IoT/event streams with batch data into one trusted table. Pick the two "
          "verticals closest to your audience and go deeper. The common thread: in every case, Delta's "
          "reliability is what makes the use case production-grade rather than a fragile experiment.")


def benefits(prs):
    s = slide(prs)
    header(s, "Platform & Value", "The Delta Lake Payoff")
    progress(s, 3)
    nodes = [
        dict(title="Reliability", icon="check", color=GREEN, desc="ACID & quality"),
        dict(title="Faster Analytics", icon="gauge", color=BLUE, desc="OPTIMIZE · Photon"),
        dict(title="Scalability", icon="cluster", color=TEAL, desc="Petabyte-scale"),
        dict(title="Governance", icon="lock", color=PURPLE, desc="Unity Catalog"),
        dict(title="Lower Cost", icon="coins", color=LAVA, desc="One open copy"),
        dict(title="Data Quality", icon="shield", color=YELLOW, desc="Enforced schema"),
    ]
    radial_infographic(s, SW / 2, 4.4, 2.05, "DELTA\nVALUE", nodes,
                       center_color=LAVA, center_d=1.75, node_d=1.15, label_size=12)
    footer(s)
    notes(s,
          "Summarize the value as a benefits wheel — the executive takeaway. RELIABILITY (ACID and quality) is "
          "the headline. FASTER ANALYTICS comes from OPTIMIZE, Z-ORDER and Photon. SCALABILITY to petabytes on "
          "cheap storage. GOVERNANCE through Unity Catalog. LOWER COST because you keep one open copy of data "
          "instead of duplicating into a separate warehouse. And DATA QUALITY from schema enforcement. Each "
          "spoke maps back to something we covered today, so this slide doubles as a recap. If a leader "
          "remembers one thing: Delta makes the data lake trustworthy enough to run the business on.")


def best_practices(prs):
    s = slide(prs)
    header(s, "Platform & Value", "Delta Lake Best Practices")
    progress(s, 3)
    checklist(s, MX + 0.2, 2.2, CW - 0.4, [
        "Use Delta as the default format for every table",
        "Follow the Bronze → Silver → Gold medallion pattern",
        "Run OPTIMIZE and Z-ORDER on large, frequently-queried tables",
        "Enable schema enforcement; evolve schemas intentionally",
        "Use MERGE for upserts and change-data-capture",
        "Partition thoughtfully — avoid too many tiny partitions",
        "Govern every table with Unity Catalog from day one",
        "Use time travel for audits, recovery and reproducible ML",
    ], cols=2, row_h=0.92, size=12.5)
    footer(s)
    notes(s,
          "Leave the audience with an actionable checklist. The most important habit: make Delta the default "
          "for every table — there's rarely a reason to use plain Parquet. Structure pipelines with the "
          "medallion pattern. Maintain performance with OPTIMIZE and Z-ORDER on big tables. Keep schema "
          "enforcement on and evolve deliberately with mergeSchema. Use MERGE for upserts and CDC rather than "
          "delete-and-reload. Be thoughtful with partitioning — over-partitioning recreates the small-file "
          "problem. Govern with Unity Catalog from the start, not as an afterthought. And lean on time travel "
          "for audits, recovery and reproducible ML. These eight habits cover the vast majority of "
          "production Delta success.")


def takeaways(prs):
    s = slide(prs)
    header(s, "Wrap-Up", "Key Takeaways")
    cols = [
        dict(no="1", color=LAVA, title="The Problem", points=[
            "Raw lakes have no reliability",
            "Bad data is costly & risky"]),
        dict(no="2", color=BLUE, title="Delta Lake", points=[
            "Open layer: Parquet + a log",
            "ACID reliability on storage"]),
        dict(no="3", color=GREEN, title="How It Works", points=[
            "Transaction log = the brain",
            "Time travel & schema control"]),
        dict(no="4", color=PURPLE, title="The Value", points=[
            "Foundation of the Lakehouse",
            "Fast, governed, trusted data"]),
    ]
    cw = (CW - 3 * 0.3) / 4
    for i, c in enumerate(cols):
        x = MX + i * (cw + 0.3)
        card(s, x, 2.0, cw, 2.55, fill=WHITE, line=LINE, shadow=True)
        band(s, x, 2.0, cw, 0.12, c["color"])
        number_badge(s, x + cw / 2, 2.62, 0.62, c["no"], fill=c["color"], size=17)
        text(s, x + 0.1, 3.05, cw - 0.2, 0.4, c["title"], size=12.5, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
        bullets(s, x + 0.26, 3.5, cw - 0.5, 1.0, c["points"], size=10.3, marker=c["color"], gap=7)
    hw = (CW - 0.3) / 2
    takeaway(s, MX, 4.95, hw, "Delta Lake turns a folder of files into a reliable, governed database.",
             color=LAVA, icon_kind="delta")
    takeaway(s, MX + hw + 0.3, 4.95, hw,
             "The transaction log delivers ACID, time travel and audit — all at once.",
             color=BLUE, icon_kind="transaction")
    takeaway(s, MX, 5.85, hw, "Schema enforcement + evolution keep tables clean yet adaptable.",
             color=GREEN, icon_kind="shield")
    takeaway(s, MX + hw + 0.3, 5.85, hw,
             "Delta is the open foundation of the Lakehouse — fast, governed and trusted.",
             color=PURPLE, icon_kind="star")
    footer(s)
    notes(s,
          "Recap the four-part journey: the PROBLEM (raw lakes are unreliable and that costs the business), "
          "DELTA LAKE (an open Parquet-plus-log layer that adds ACID), HOW IT WORKS (the transaction log is "
          "the brain, enabling time travel and schema control), and THE VALUE (Delta is the trusted, fast, "
          "governed foundation of the Lakehouse). Quiz the room on the four banner statements. If they can "
          "explain why the transaction log makes a lake reliable, the day succeeded. Transition to Q&A.")


def closing(prs):
    s = slide(prs, bg=NAVY)
    dark_bg(s)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 2.2, 0.09, 1.5, fill=LAVA, line=None, adj=0.5)
    text(s, MX + 0.3, 2.05, 10, 1.1, "Questions?", size=52, color=WHITE, bold=True)
    text(s, MX + 0.32, 3.25, 10.6, 0.6,
         "You now understand Delta Lake from the transaction log up — and why it's the foundation of every "
         "reliable Lakehouse.", size=15, color=RGBColor(0xC4, 0xD0, 0xD4), spacing=1.2)
    text(s, MX + 0.32, 4.25, 8, 0.3, "GO DEEPER", size=11.5, color=LAVA_LT, bold=True, font=FONT_SB)
    res = ["docs.delta.io", "Databricks Academy", "Free Edition — hands-on", "delta.io open source"]
    cx = MX + 0.32
    for r in res:
        w = 0.32 + len(r) * 0.105
        chip(s, cx, 4.62, w, 0.52, r, NAVY_2, WHITE, size=11, radius=0.5, line=LINE_DK)
        cx += w + 0.22
    recap = [("Reliability", "check", GREEN), ("ACID", "transaction", LAVA),
             ("Time Travel", "clock", BLUE), ("Schema", "shield", YELLOW),
             ("Medallion", "diamond", GOLD), ("Lakehouse", "warehouse", PURPLE)]
    cx = MX + 0.32
    for name, ic, col in recap:
        icon_tile(s, cx + 0.4, 5.85, 0.78, ic, col, shadow=True)
        text(s, cx - 0.1, 6.34, 1.0, 0.4, name, size=9.5, color=RGBColor(0xC4, 0xD0, 0xD4),
             align=PP_ALIGN.CENTER, spacing=0.9)
        cx += 1.15
    text(s, 8.7, 6.55, CW + MX - 8.7, 0.3, "Next: hands-on Delta Lake lab",
         size=12, color=GRAY_LT, italic=True, align=PP_ALIGN.RIGHT)
    notes(s,
          "Close and open the floor. Recap the six icons — reliability, ACID, time travel, schema control, "
          "medallion and the Lakehouse. Point to resources: docs.delta.io and delta.io (Delta is open "
          "source — anyone can use it beyond Databricks), Databricks Academy for structured learning, and the "
          "free edition to practice. Preview the next session: a hands-on lab creating Delta tables, running "
          "MERGE and OPTIMIZE, and demonstrating time travel live. Invite questions.")


# ===========================================================================
def main():
    prs = new_deck()
    prs.core_properties.title = "Databricks Day 3 — Delta Lake"
    prs.core_properties.subject = "Delta Lake: reliability for the data lake"

    cover(prs)
    agenda(prs)
    s_divider(prs, "01", "Part One", "The Problem",
              ["Why traditional data lakes fail", "The business cost of poor data quality"],
              "warning", color=LAVA)
    why_lakes_fail(prs)
    business_impact(prs)
    s_divider(prs, "02", "Part Two", "Delta Lake Fundamentals",
              ["What Delta Lake is", "Architecture & the five core components"],
              "delta", color=LAVA)
    what_is_delta(prs)
    delta_architecture(prs)
    delta_components(prs)
    s_divider(prs, "03", "Part Three", "How Delta Lake Works",
              ["ACID, the transaction log & reliability", "Time travel, schema control & streaming"],
              "transaction", color=LAVA)
    acid(prs)
    transaction_log(prs)
    reliability_flow(prs)
    time_travel(prs)
    schema_enforcement(prs)
    schema_evolution(prs)
    batch_streaming(prs)
    s_divider(prs, "04", "Part Four", "Platform & Value",
              ["Medallion, Lakehouse, performance & governance", "End-to-end flow, use cases & best practices"],
              "gauge", color=LAVA)
    medallion(prs)
    lakehouse(prs)
    performance(prs)
    governance(prs)
    end_to_end(prs)
    use_cases(prs)
    benefits(prs)
    best_practices(prs)
    takeaways(prs)
    closing(prs)

    prs.save(OUT)
    print(f"Saved {OUT} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    main()
