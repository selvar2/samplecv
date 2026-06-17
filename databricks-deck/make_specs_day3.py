#!/usr/bin/env python3
"""make_specs_day3.py — per-slide design doc & facilitator guide for the Delta Lake deck.
Speaker notes are read from the built .pptx so the doc always matches the deck."""
from pptx import Presentation

PPTX = "Databricks-Day3-Delta-Lake.pptx"
OUT = "SLIDE-SPECIFICATIONS-Day3.md"

SPECS = [
    dict(title="Title / Cover", objective="Open Day 3 and frame the Delta Lake journey.",
         content="Title, subtitle, four part-chips, audience and positioning.",
         visual="Dark navy hero with a large Delta-triangle motif and lava accents.",
         diagram="—", design="Navy + lava; bold title; Delta triangle as the hero mark."),
    dict(title="Day 3 Roadmap — What You'll Learn", objective="Preview the four parts and their outcomes.",
         content="Four part cards (Problem, Fundamentals, How It Works, Platform & Value).",
         visual="Four icon cards with topics and an outcome line each.",
         diagram="Four-column roadmap.", design="Color-code parts; one outcome per card."),
    dict(title="Part 01 Divider — The Problem", objective="Transition into why lakes fail.",
         content="Part number, title, themes.", visual="Navy divider, giant '01', warning icon.",
         diagram="—", design="Shared divider system across four parts."),
    dict(title="Why Traditional Data Lakes Fail", objective="Establish the core problem the course solves.",
         content="A 'data swamp' banner plus five failure modes.",
         visual="Navy banner (lake + warning) over five red-themed problem cards.",
         diagram="Broken data-lake infographic.",
         design="Red/orange problem theme; vivid 'data swamp' framing."),
    dict(title="The Business Cost of Bad Data", objective="Translate technical problems into business risk.",
         content="Four impact cards plus a stat panel.",
         visual="2x2 impact cards beside a navy statistics panel.",
         diagram="—", design="Lead with money & risk; present figures as industry estimates."),
    dict(title="Part 02 Divider — Delta Lake Fundamentals", objective="Transition into what Delta is.",
         content="Part number, title, themes.", visual="Navy divider, giant '02', Delta icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="What Is Delta Lake?", objective="Define Delta Lake and preview its capabilities.",
         content="Radial feature wheel plus a definition card and bullets.",
         visual="Six-node radial 'wheel' on the left; navy definition card on the right.",
         diagram="Radial infographic: 6 features around a DELTA LAKE hub.",
         design="Stress open format, unified batch+stream, Lakehouse foundation."),
    dict(title="Delta Lake Architecture", objective="Show where Delta sits and what a table contains.",
         content="Users→Databricks→Delta→Storage layers; the three parts inside a Delta table.",
         visual="Left layered stack with down-arrows; right 'inside a table' panel.",
         diagram="Users ↓ Databricks ↓ Delta Lake ↓ Cloud Storage  +  log/metadata/data files.",
         design="The transaction log is teased as the 'brain'."),
    dict(title="The Five Core Components", objective="Summarize Delta's moving parts.",
         content="Delta Table, Transaction Log, Metadata, Storage Layer, Optimization.",
         visual="Five-node radial wheel around a DELTA LAKE hub.",
         diagram="Radial infographic (5 components).", design="One color per component."),
    dict(title="Part 03 Divider — How Delta Lake Works", objective="Transition into the internals.",
         content="Part number, title, themes.", visual="Navy divider, giant '03', transaction icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="ACID Transactions Explained", objective="Teach the four ACID guarantees.",
         content="Atomicity, Consistency, Isolation, Durability with examples.",
         visual="Four premium letter-cards (A/C/I/D) with icons + a note band.",
         diagram="Four-card ACID infographic.", design="Big letter badges; one color per property."),
    dict(title="The Transaction Log (_delta_log)", objective="Explain the log — the source of every guarantee.",
         content="Write request → log update → atomic commit → read consistency; versioned commits.",
         visual="Four-step carousel over a versioned _delta_log strip.",
         diagram="Write → Log Update → Commit → Read Consistency (00000.json … 00003.json).",
         design="Frame the log as the table's 'brain'."),
    dict(title="How Delta Ensures Reliability", objective="Show the safe path every write follows.",
         content="Write → Validate → Commit → Version → Read.",
         visual="Five-step carousel with arrows and a note band.",
         diagram="Write → Validate → Commit → Version → Read.",
         design="Reliability by design, not by hope."),
    dict(title="Time Travel & Data Versioning", objective="Show versioning and instant rollback.",
         content="v0→v1→v2→v3 with a bad update rolled back; restore callout.",
         visual="Version timeline cards with a rollback callout and note band.",
         diagram="v0 → v1 → v2 (bad) → v3 (rollback to v1).",
         design="Mistakes become undoable — audits & reproducible ML."),
    dict(title="Schema Enforcement", objective="Show how Delta blocks accidental corruption.",
         content="Without Delta vs With Delta.",
         visual="Two contrasting panels separated by a VS badge.",
         diagram="Before/after comparison.", design="Red 'without' vs green 'with'; gatekeeper framing."),
    dict(title="Schema Evolution", objective="Show how schemas change safely and intentionally.",
         content="New Column → Validation → Evolution → Updated Table.",
         visual="Four-step carousel with a note band.",
         diagram="New Column → Validation → Evolution (mergeSchema) → Updated Table.",
         design="Enforcement blocks accidental; evolution allows intentional."),
    dict(title="Unified Batch + Streaming", objective="Show one table serving batch and streaming.",
         content="Batch & streaming sources merge into one Delta table → BI/ML/Analytics.",
         visual="Split sources on the left converge on a central Delta hub feeding consumers.",
         diagram="Batch + Streaming → Delta Lake (one table) → BI / ML / Analytics.",
         design="Eliminates the two-pipeline (Lambda) problem."),
    dict(title="Part 04 Divider — Platform & Value", objective="Transition into platform applications.",
         content="Part number, title, themes.", visual="Navy divider, giant '04', gauge icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="Delta Powers the Medallion Architecture", objective="Connect Delta to the medallion pattern.",
         content="Bronze, Silver, Gold — each an ACID Delta table.",
         visual="Three detail cards with arrows and a bronze→gold quality gradient.",
         diagram="Bronze → Silver → Gold.", design="Reliability compounds because every layer is Delta."),
    dict(title="Delta Lake + the Lakehouse", objective="Show Delta as the keystone of the Lakehouse.",
         content="Cloud Storage + Delta Lake + Databricks = Lakehouse.",
         visual="Three input cards joined by +/= to a navy LAKEHOUSE result card.",
         diagram="Storage + Delta + Databricks = Lakehouse.",
         design="Without Delta you just have a lake."),
    dict(title="Performance Optimization", objective="Show that Delta is fast, not just reliable.",
         content="OPTIMIZE, Z-ORDER, Compaction, Caching + the payoff.",
         visual="2x2 technique cards beside a navy 'payoff' gauge panel.",
         diagram="—", design="Data skipping + fewer files + Photon = faster, cheaper."),
    dict(title="Governance with Unity Catalog", objective="Pair reliability with governance.",
         content="Unity Catalog → Security → Audit & Lineage → Compliance.",
         visual="Four-step carousel with a note band.",
         diagram="Unity Catalog → Security → Audit/Lineage → Compliance.",
         design="Delta makes data trustworthy; UC controls & proves its use."),
    dict(title="End-to-End Data Flow on Delta", objective="Tie everything into one pipeline (capstone).",
         content="Sources → Ingestion → Bronze → Silver → Gold → Dashboard → AI/ML on Delta.",
         visual="Seven numbered stages over a Delta Lake foundation band.",
         diagram="Sources → Ingestion → Bronze → Silver → Gold → Dashboard → AI/ML.",
         design="Every stage is a Delta table — one open copy end to end."),
    dict(title="Delta Lake in the Real World", objective="Show industry breadth.",
         content="Six verticals, each tied to a Delta capability.",
         visual="2x3 industry card grid.", diagram="—",
         design="Anchor each card to a feature learned today."),
    dict(title="The Delta Lake Payoff", objective="Summarize the value as a benefits wheel.",
         content="Reliability, Faster Analytics, Scalability, Governance, Lower Cost, Data Quality.",
         visual="Six-node radial benefits wheel.",
         diagram="Radial infographic (6 benefits).", design="Doubles as a recap; each spoke maps to a topic."),
    dict(title="Delta Lake Best Practices", objective="Leave an actionable checklist.",
         content="Eight production best practices.",
         visual="Two-column checklist with green check badges.",
         diagram="Checklist infographic.", design="Default to Delta; medallion; OPTIMIZE; govern early."),
    dict(title="Key Takeaways", objective="Recap the four parts and land four takeaways.",
         content="Four recap cards + four takeaway banners.",
         visual="Four numbered recap cards over four icon banners.",
         diagram="—", design="Quiz the room on the banners."),
    dict(title="Questions? / Closing", objective="Close and point to resources.",
         content="Closing, resource chips, six-icon recap, next-session preview.",
         visual="Navy closing slide with resource chips and recap icons.",
         diagram="—", design="Mirror the cover; preview the hands-on lab."),
]


def main():
    prs = Presentation(PPTX)
    slides = list(prs.slides)
    assert len(slides) == len(SPECS), f"{len(slides)} vs {len(SPECS)}"
    L = []
    A = L.append
    A("# Databricks Day 3 — Delta Lake — Slide Specifications & Facilitator Guide\n")
    A("_Reliability for the Data Lake · Enterprise Data Platform Training Series_\n")
    A("Specifies every slide in **Databricks-Day3-Delta-Lake.pptx**. Speaker notes are "
      "extracted directly from the deck, so they always match the notes pane.\n")
    A("> **Format:** Objective · Content · Visual Layout · Diagram · Speaker Notes · Design Tips\n")
    A("\n---\n")
    breaks = {3: "Part 1 · The Problem", 6: "Part 2 · Delta Lake Fundamentals",
              10: "Part 3 · How Delta Lake Works", 18: "Part 4 · Platform & Value"}
    for i, (sp, sl) in enumerate(zip(SPECS, slides), start=1):
        if i in breaks:
            A(f"\n## {breaks[i]}\n")
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
        A(f"### Slide {i}: {sp['title']}\n")
        A(f"**Objective:** {sp['objective']}\n")
        A(f"**Content:** {sp['content']}\n")
        A(f"**Visual Layout:** {sp['visual']}\n")
        A(f"**Diagram:** {sp['diagram']}\n")
        A(f"**Speaker Notes:** {notes}\n")
        A(f"**Design Tips:** {sp['design']}\n")
        A("\n---\n")
    open(OUT, "w").write("\n".join(L))
    print(f"Wrote {OUT} ({len(SPECS)} slides)")


if __name__ == "__main__":
    main()
