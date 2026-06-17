#!/usr/bin/env python3
"""make_specs_day4.py — per-slide design doc & facilitator guide for the Data Processing deck.
Speaker notes are read from the built .pptx so the doc always matches the deck."""
from pptx import Presentation

PPTX = "Databricks-Day4-Data-Processing-Concepts.pptx"
OUT = "SLIDE-SPECIFICATIONS-Day4.md"

SPECS = [
    dict(title="Title / Cover", objective="Open Day 4 and frame the data-processing journey.",
         content="Title, subtitle, five part-chips, audience and positioning.",
         visual="Dark navy hero with a 'data flowing through a pipeline' motif.",
         diagram="—", design="Navy + lava; vendor-neutral concepts in a Databricks context."),
    dict(title="Why Data Processing Matters", objective="Establish why processing creates value.",
         content="Raw → Insights → Value flow plus four business outcomes.",
         visual="Three-step value flow over a 2x... outcome card row.",
         diagram="Raw Data → Insights → Business Value.",
         design="Lead with value (decisions, automation, innovation, CX)."),
    dict(title="The Data Journey — Seven Stages", objective="Give the spine of the whole deck.",
         content="Generate → Collect → Process → Store → Analyze → Visualize → Act.",
         visual="Seven numbered icon-stage cards with arrows + note band.",
         diagram="Generate → Collect → Process → Store → Analyze → Visualize → Act.",
         design="This is the map; each section zooms into a stage."),
    dict(title="Part 01 Divider — Data Sources & Types", objective="Transition into sources & types.",
         content="Part number, title, themes.", visual="Navy divider, giant '01', globe icon.",
         diagram="—", design="Shared divider system across five parts."),
    dict(title="Where Enterprise Data Comes From", objective="Survey the variety of data sources.",
         content="Databases, APIs, IoT, Applications, Logs/Files, Social/Web.",
         visual="2x3 icon card grid.", diagram="—",
         design="Variety in structure/speed/volume motivates ingestion."),
    dict(title="Structured vs. Semi vs. Unstructured", objective="Classify data by structure.",
         content="Examples, schema, storage and share for each type.",
         visual="Three comparison cards with mini-tables.",
         diagram="3-column data-type comparison.",
         design="~80% is unstructured — why warehouses alone fall short."),
    dict(title="Part 02 Divider — Data Ingestion", objective="Transition into ingestion.",
         content="Part number, title, themes.", visual="Navy divider, giant '02', funnel icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="What Is Data Ingestion?", objective="Define ingestion as step one.",
         content="Sources → Ingestion Layer → Storage.",
         visual="Three-stage labelled pipeline with example cards + note band.",
         diagram="Sources → Ingestion Layer → Storage.",
         design="Reliability (no loss/dupes) is the hard part."),
    dict(title="Batch Ingestion", objective="Explain scheduled, chunked loading.",
         content="Nightly/periodic loads on a clock.",
         visual="Clock-style timeline of scheduled runs + note band.",
         diagram="Timeline of periodic loads (02:00 … 18:00).",
         design="Simple, cheap, ideal for history; latency is the trade-off."),
    dict(title="Streaming Ingestion", objective="Explain continuous, real-time ingestion.",
         content="Kafka/IoT/events → continuous ingest → real-time Delta.",
         visual="Sources converge on a continuous-ingest hub feeding a Delta table.",
         diagram="Kafka · IoT · Events → continuous ingest → Delta Bronze.",
         design="For fraud, monitoring, live dashboards; Structured Streaming."),
    dict(title="Batch vs. Streaming", objective="Help choose the right ingestion mode.",
         content="Timing, latency, size, complexity, cost, best-for.",
         visual="Three-column comparison table.",
         diagram="Comparison table (Aspect | Batch | Streaming).",
         design="Not either/or — Delta lets one table serve both."),
    dict(title="Part 03 Divider — Transformation · ETL & ELT", objective="Transition into transformation.",
         content="Part number, title, themes.", visual="Navy divider, giant '03', gear icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="What Is Data Transformation?", objective="Define transformation as the core of processing.",
         content="Raw → Clean → Standardize → Enrich → Processed.",
         visual="Five-step carousel + note band.",
         diagram="Raw → Clean → Standardize → Enrich → Processed.",
         design="Maps directly onto Bronze → Silver → Gold."),
    dict(title="Common Transformations", objective="Show the building-block operations.",
         content="Filter, Aggregate, Join, Sort, Deduplicate, Validate.",
         visual="2x3 icon card grid.", diagram="—",
         design="Most pipelines compose these six in Spark/SQL."),
    dict(title="ETL — Extract, Transform, Load", objective="Explain the classic warehouse pattern.",
         content="Extract → Transform (separate engine) → Load.",
         visual="Three-step carousel + note band.",
         diagram="Extract → Transform → Load.",
         design="Transform BEFORE load; only clean data lands."),
    dict(title="ELT — Extract, Load, Transform", objective="Explain the modern cloud pattern.",
         content="Extract → Load raw → Transform in-place.",
         visual="Three-step carousel + note band.",
         diagram="Extract → Load → Transform.",
         design="Load first; cheap storage + elastic compute. The medallion model."),
    dict(title="ETL vs. ELT", objective="Contrast the two and explain the shift.",
         content="Order, where transform happens, raw kept, scale, cost, cloud-fit.",
         visual="Three-column comparison table.",
         diagram="Comparison table (Aspect | ETL | ELT).",
         design="ELT is the modern Lakehouse standard."),
    dict(title="Part 04 Divider — Distributed Processing & Pipelines", objective="Transition into processing.",
         content="Part number, title, themes.", visual="Navy divider, giant '04', cluster icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="Why Distributed Processing?", objective="Explain scale-out vs scale-up.",
         content="Single server vs cluster processing.",
         visual="Before/after panels separated by a VS badge.",
         diagram="Single Server vs Cluster Processing.",
         design="Big data outgrew the single machine — scale out."),
    dict(title="The Apache Spark Processing Model", objective="Introduce Spark's execution model.",
         content="Driver → Executors → Tasks.",
         visual="Three-step carousel + note band.",
         diagram="Driver → Executors → Tasks.",
         design="Driver coordinates, executors work, tasks run in parallel."),
    dict(title="How Spark Processes Data", objective="Show the data's-eye view of a job.",
         content="Read → Partition → Transform → Execute → Write.",
         visual="Five-step carousel + note band.",
         diagram="Read → Partition → Transform → Execute → Write.",
         design="Partitioning is the secret to scale; lazy + in-memory."),
    dict(title="What Is a Data Pipeline?", objective="Define the automated data flow.",
         content="Ingest → Process → Store → Serve.",
         visual="Four-stage labelled pipeline + note band.",
         diagram="Ingest → Process → Store → Serve.",
         design="Automated, repeatable, reliable — a factory line for data."),
    dict(title="Modern Data Pipeline Architecture", objective="Give the reference pipeline blueprint.",
         content="Sources → Ingestion → Processing → Storage → Analytics + governance.",
         visual="Five numbered stages over a governance/orchestration band.",
         diagram="Sources → Ingestion → Processing → Storage → Analytics.",
         design="Maps to Auto Loader, Spark, Delta, BI/ML on Databricks."),
    dict(title="Pipeline Orchestration", objective="Explain what makes pipelines production-grade.",
         content="Jobs & Tasks, Scheduling, Dependencies, Monitoring.",
         visual="2x2 icon card grid.", diagram="—",
         design="Orchestration turns scripts into reliable pipelines (Workflows)."),
    dict(title="Part 05 Divider — Quality, Storage & Databricks", objective="Transition into quality & storage.",
         content="Part number, title, themes.", visual="Navy divider, giant '05', shield icon.",
         diagram="—", design="Consistent divider system."),
    dict(title="Why Data Quality Matters", objective="Show the cost of poor quality.",
         content="Costly mistakes, broken trust, wasted time, compliance risk + the principle.",
         visual="2x2 impact cards beside a 'garbage in, garbage out' principle panel.",
         diagram="—", design="Processing can't invent correctness — quality is upstream."),
    dict(title="The Six Dimensions of Data Quality", objective="Make quality measurable.",
         content="Accuracy, Completeness, Consistency, Timeliness, Validity, Uniqueness.",
         visual="Six-node radial 'wheel' infographic.",
         diagram="Radial infographic (6 quality dimensions).",
         design="Delta/DLT expectations automate many of these checks."),
    dict(title="Warehouse vs. Lake vs. Lakehouse", objective="Compare the storage paradigms.",
         content="Strengths & limits of each; Lakehouse as best of both.",
         visual="Three comparison cards; Lakehouse highlighted with a 'best of both' chip.",
         diagram="3-way storage comparison.",
         design="The Lakehouse removes the warehouse-vs-lake choice."),
    dict(title="Medallion Architecture", objective="Show where transformation, quality & storage meet.",
         content="Bronze, Silver, Gold with rising quality.",
         visual="Three detail cards with arrows and a bronze→gold gradient.",
         diagram="Bronze → Silver → Gold.",
         design="The default shape of a data platform; every layer is Delta."),
    dict(title="End-to-End Databricks Processing", objective="Realize every concept on Databricks (capstone).",
         content="Sources → Auto Loader → Delta → Spark → Lakehouse → BI/AI.",
         visual="Six numbered stages over a Unity Catalog / Workflows band.",
         diagram="Sources → Auto Loader → Delta Lake → Spark → Lakehouse → BI/AI.",
         design="Map each box back to a concept covered today."),
    dict(title="Data Processing Best Practices", objective="Leave an actionable checklist.",
         content="Eight production best practices.",
         visual="Two-column checklist with green check badges.",
         diagram="Checklist infographic.",
         design="ELT, medallion, quality-everywhere, idempotency, orchestration."),
    dict(title="Data Processing in the Real World", objective="Prove the concepts across industries.",
         content="Retail, Banking, Telecom, Healthcare, Manufacturing, E-Commerce.",
         visual="2x3 industry card grid.", diagram="—",
         design="Trace the full journey for the audience's two closest verticals."),
    dict(title="Key Takeaways", objective="Recap the journey and land four takeaways.",
         content="Four recap cards + four takeaway banners.",
         visual="Four numbered recap cards over four icon banners.",
         diagram="—", design="Quiz the room on the banners."),
    dict(title="Questions? / Closing", objective="Close and point to resources.",
         content="Closing, resource chips, six-icon recap, next-session preview.",
         visual="Navy closing slide with resource chips and recap icons.",
         diagram="—", design="Mirror the cover; preview the hands-on pipeline lab."),
]


def main():
    prs = Presentation(PPTX)
    slides = list(prs.slides)
    assert len(slides) == len(SPECS), f"{len(slides)} vs {len(SPECS)}"
    L = []
    A = L.append
    A("# Databricks Day 4 — Data Processing Concepts — Slide Specifications & Facilitator Guide\n")
    A("_From raw signals to enterprise insight · Enterprise Data Platform Training Series_\n")
    A("Specifies every slide in **Databricks-Day4-Data-Processing-Concepts.pptx**. Speaker notes "
      "are extracted directly from the deck, so they always match the notes pane.\n")
    A("> **Format:** Objective · Content · Visual Layout · Diagram · Speaker Notes · Design Tips\n")
    A("\n---\n")
    breaks = {4: "Part 1 · Data Sources & Types", 7: "Part 2 · Data Ingestion",
              12: "Part 3 · Transformation · ETL & ELT",
              18: "Part 4 · Distributed Processing & Pipelines",
              25: "Part 5 · Quality, Storage & Databricks"}
    for i, (sp, sl) in enumerate(zip(SPECS, slides), start=1):
        if i in breaks:
            A(f"\n## {breaks[i]}\n")
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
        A(f"### Slide {i}: {sp['title']}\n")
        A(f"**Objective:** {sp['objective']}\n")
        A(f"**Content:** {sp['content']}\n")
        A(f"**Visual Layout:** {sp['visual']}\n")
        A(f"**Diagram:** {sp['diagram']}\n")
        A(f"**Speaker Notes:** {notes}\n")
        A(f"**Design Tips:** {sp['design']}\n")
        A("\n---\n")
    open(OUT, "w").write("\n".join(L))
    print(f"Wrote {OUT} ({len(SPECS)} slides)")


if __name__ == "__main__":
    main()
